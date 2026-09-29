from __future__ import annotations

from calendar import monthrange
from collections import OrderedDict
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from html import escape as html_escape
from io import BytesIO
import re
from typing import Any, Iterable

from django.core.exceptions import PermissionDenied, ValidationError
from django.db.models import Q
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.accounts.access_catalog import AccessPermission
from apps.accounts.access_policy import membership_allows_project, membership_has_permission
from apps.rental_manpower.models import (
    RentalTimesheetEntry,
    RentalTimesheetOvertime,
    RentalTimesheetPeriod,
    WorkerAssignment,
)
from apps.rental_manpower.project_adapter import rental_project_for_company
from apps.rental_manpower.selectors.timesheets import _project_worker_queryset
from apps.rental_manpower.services.timesheets import month_bounds
from apps.rental_manpower.services.timesheet_policy import (
    effective_timesheet_policy, overtime_bill_rate, split_daily_hours,
)


EXPORT_MAX_ROWS = 5000
VALID_EXPORT_FORMATS = {"xlsx", "pdf"}
VALID_EXPORT_SCOPES = {"all_matching", "current_page", "selected"}


@dataclass(frozen=True)
class ExportColumn:
    key: str
    label: str
    group: str
    width: int = 16
    pdf_width_mm: float = 24.0
    kind: str = "text"
    default: bool = False
    permission: str = "base"

    def payload(self) -> dict[str, Any]:
        return {
            "key": self.key,
            "label": self.label,
            "group": self.group,
            "width": self.width,
            "kind": self.kind,
            "default": self.default,
        }


_STATIC_COLUMNS: tuple[ExportColumn, ...] = (
    ExportColumn("worker_id", "Worker ID", "Worker", 15, 22, default=True),
    ExportColumn("worker_name", "Worker name", "Worker", 28, 38, default=True),
    ExportColumn("national_id", "National ID / Iqama", "Worker", 20, 28, permission="worker_detail"),
    ExportColumn("worker_phone", "Worker phone", "Worker", 18, 26, permission="worker_detail"),
    ExportColumn("worker_status", "Worker status", "Worker", 14, 22, permission="worker_detail"),
    ExportColumn("worker_notes", "Worker notes", "Worker", 32, 44, permission="worker_detail"),

    ExportColumn("supplier_code", "Supplier code", "Supplier", 16, 24),
    ExportColumn("supplier_name", "Supplier name", "Supplier", 28, 38, default=True),
    ExportColumn("supplier_contact", "Supplier contact", "Supplier", 24, 34, permission="supplier_detail"),
    ExportColumn("supplier_phone", "Supplier phone", "Supplier", 18, 27, permission="supplier_detail"),
    ExportColumn("supplier_email", "Supplier email", "Supplier", 28, 40, permission="supplier_detail"),
    ExportColumn("supplier_cr", "Supplier CR", "Supplier", 18, 27, permission="supplier_detail"),
    ExportColumn("supplier_vat", "Supplier VAT", "Supplier", 20, 30, permission="supplier_detail"),
    ExportColumn("supplier_payment_terms", "Payment terms", "Supplier", 22, 32, permission="supplier_detail"),
    ExportColumn("supplier_address", "Supplier address", "Supplier", 36, 48, permission="supplier_detail"),

    ExportColumn("project_code", "Project code", "Project", 16, 24),
    ExportColumn("project_name", "Project name", "Project", 28, 38, default=True),
    ExportColumn("project_client", "Client / principal", "Project", 26, 36),
    ExportColumn("project_location", "Project location", "Project", 24, 34),
    ExportColumn("project_manager", "Project manager", "Project", 22, 32),
    ExportColumn("period", "Period", "Project", 15, 22, default=True),
    ExportColumn("regular_hours_limit", "Regular hours / day", "Project", 18, 26, kind="number"),

    ExportColumn("assignment_trade", "Trade / role", "Assignment", 24, 34, default=True),
    ExportColumn("assignment_rate_type", "Rate type", "Assignment", 16, 24),
    ExportColumn("assignment_rate", "Commercial rate", "Assignment", 18, 27, permission="commercial"),
    ExportColumn("assignment_start", "Assignment start", "Assignment", 16, 24),
    ExportColumn("assignment_end", "Assignment end", "Assignment", 16, 24),
    ExportColumn("assignment_segments", "Assignment segments", "Assignment", 42, 58),

    ExportColumn("assigned_days", "Assigned days", "Totals & exceptions", 14, 20, kind="number"),
    ExportColumn("regular_hours", "Regular hours", "Totals & exceptions", 15, 22, kind="number", default=True),
    ExportColumn("total_hours", "Total hours", "Totals & exceptions", 15, 22, kind="number", default=True),
    ExportColumn("work_days", "Worked days", "Totals & exceptions", 14, 20, kind="number"),
    ExportColumn("absent_days", "Absent days", "Totals & exceptions", 14, 20, kind="number"),
    ExportColumn("no_scope_days", "No-scope days", "Totals & exceptions", 15, 22, kind="number"),
    ExportColumn("leave_days", "Leave days", "Totals & exceptions", 14, 20, kind="number"),
    ExportColumn("off_days", "Off days", "Totals & exceptions", 12, 18, kind="number"),
    ExportColumn("zero_hour_days", "Zero-hour days", "Totals & exceptions", 16, 23, kind="number"),
    ExportColumn("missing_days", "Missing days", "Totals & exceptions", 14, 20, kind="number", default=True),
    ExportColumn("daily_notes", "Daily remarks", "Totals & exceptions", 40, 56),

    ExportColumn("overtime_hours", "OT hours", "Overtime", 14, 20, kind="number", default=True),
    ExportColumn("automatic_overtime_hours", "Auto OT hours", "Overtime", 16, 23, kind="number"),
    ExportColumn("additional_overtime_hours", "Additional OT hours", "Overtime", 19, 27, kind="number"),
    ExportColumn("overtime_multiplier", "OT multiplier", "Overtime", 15, 22, kind="number", permission="commercial"),
    ExportColumn("overtime_rate", "OT hourly rate", "Overtime", 16, 24, kind="number", permission="commercial"),
    ExportColumn("base_wage", "Base wage", "Overtime", 18, 27, kind="number", permission="commercial"),
    ExportColumn("overtime_wage", "OT wage", "Overtime", 18, 27, kind="number", permission="commercial"),
    ExportColumn("gross_wage", "Gross wage", "Overtime", 18, 27, kind="number", permission="commercial"),

    ExportColumn("timesheet_status", "Timesheet status", "Workflow", 17, 25, default=True),
    ExportColumn("timesheet_revision", "Revision", "Workflow", 11, 18, kind="number"),
    ExportColumn("submitted_at", "Submitted at", "Workflow", 21, 31, permission="workflow_detail"),
    ExportColumn("submitted_by", "Submitted by", "Workflow", 22, 32, permission="workflow_detail"),
    ExportColumn("approved_at", "Approved at", "Workflow", 21, 31, permission="workflow_detail"),
    ExportColumn("approved_by", "Approved by", "Workflow", 22, 32, permission="workflow_detail"),
    ExportColumn("locked_at", "Locked at", "Workflow", 21, 31, permission="workflow_detail"),
    ExportColumn("locked_by", "Locked by", "Workflow", 22, 32, permission="workflow_detail"),
)


def _can(membership, permission: AccessPermission) -> bool:
    return bool(membership and membership_has_permission(membership, permission))


def _permission_flags(membership) -> dict[str, bool]:
    return {
        "base": True,
        "worker_detail": _can(membership, AccessPermission.RENTAL_WORKERS_VIEW),
        "supplier_detail": _can(membership, AccessPermission.RENTAL_SUPPLIERS_VIEW),
        "commercial": _can(membership, AccessPermission.RENTAL_SETTLEMENTS_VIEW)
        or _can(membership, AccessPermission.RENTAL_ASSIGNMENTS_MANAGE),
        "workflow_detail": _can(membership, AccessPermission.RENTAL_TIMESHEETS_APPROVE)
        or _can(membership, AccessPermission.RENTAL_REPORTS_VIEW),
    }


def _day_columns(period_start: date) -> list[ExportColumn]:
    days = monthrange(period_start.year, period_start.month)[1]
    result = []
    for day in range(1, days + 1):
        current = period_start.replace(day=day)
        result.append(
            ExportColumn(
                f"day_{day:02d}",
                f"{day:02d} {current:%a}",
                "Daily entries",
                width=11,
                pdf_width_mm=13,
                default=True,
            )
        )
    return result


def _available_columns(period_start: date, membership) -> list[ExportColumn]:
    flags = _permission_flags(membership)
    columns: list[ExportColumn] = []
    inserted_daily = False
    for spec in _STATIC_COLUMNS:
        if spec.group == "Totals & exceptions" and not inserted_daily:
            columns.extend(_day_columns(period_start))
            inserted_daily = True
        if flags.get(spec.permission, False):
            columns.append(spec)
    return columns


def _normalized_supplier_id(value: object) -> str:
    raw = str(value or "").strip()
    return "" if raw.lower() in {"", "all", "all suppliers"} else raw


def _resolve_project(*, company, membership, project_id):
    project = rental_project_for_company(company=company, identifier=project_id)
    if not membership_allows_project(membership, project):
        raise PermissionDenied("This Rental project is outside your assigned access scope.")
    return project


def rental_timesheet_export_schema(
    *, company, membership, project_id, period_start: date, query: str = "", supplier_id: str = ""
) -> dict[str, Any]:
    start, end = month_bounds(period_start)
    project = _resolve_project(company=company, membership=membership, project_id=project_id)
    supplier_id = _normalized_supplier_id(supplier_id)
    matching_count = _project_worker_queryset(
        company=company,
        project=project,
        start=start,
        end=end,
        query=(query or "").strip(),
        supplier_id=supplier_id,
    ).count()
    period = (
        RentalTimesheetPeriod.objects.for_company(company)
        .filter(project=project, period_start=start)
        .first()
    )
    columns = _available_columns(start, membership)
    grouped: OrderedDict[str, list[dict[str, Any]]] = OrderedDict()
    for column in columns:
        grouped.setdefault(column.group, []).append(column.payload())
    return {
        "project": {"id": str(project.reference), "code": project.code, "name": project.name},
        "period": {
            "value": f"{start:%Y-%m}",
            "label": f"{start:%B %Y}",
            "status": period.get_status_display() if period else "Draft",
            "revision": period.revision if period else 0,
        },
        "filters": {"query": (query or "").strip(), "supplierId": supplier_id},
        "matchingCount": matching_count,
        "maxRows": EXPORT_MAX_ROWS,
        "formats": [
            {"value": "xlsx", "label": "Excel (.xlsx)", "description": "Full-detail workbook with filters, frozen headers and selected columns."},
            {"value": "pdf", "label": "PDF (.pdf)", "description": "Shareable landscape report; wide selections are split into readable column bands."},
        ],
        "groups": [{"label": label, "columns": items} for label, items in grouped.items()],
        "defaultColumns": [column.key for column in columns if column.default],
        "allColumns": [column.key for column in columns],
    }


def _actor_label(user) -> str:
    if user is None:
        return ""
    return (user.get_full_name() or "").strip() or user.username


def _date_time_label(value) -> str:
    if not value:
        return ""
    try:
        local = timezone.localtime(value)
    except Exception:
        local = value
    return local.strftime("%Y-%m-%d %H:%M")


def _entry_display(entry: RentalTimesheetEntry | None) -> str:
    if entry is None:
        return ""
    if entry.code:
        return entry.code
    value = entry.regular_hours
    if value == value.to_integral():
        return str(int(value))
    return format(value.normalize(), "f")


def _segment_label(assignment: WorkerAssignment) -> str:
    end = assignment.effective_to.isoformat() if assignment.effective_to else "Open"
    return f"{assignment.effective_from.isoformat()} → {end}: {assignment.trade} / {assignment.get_rate_type_display()} / {assignment.rate}"


def _select_worker_queryset(
    *,
    company,
    project,
    start: date,
    end: date,
    query: str,
    supplier_id: str,
    scope: str,
    page: int,
    page_size: int,
    worker_ids: Iterable[str],
):
    base = _project_worker_queryset(
        company=company,
        project=project,
        start=start,
        end=end,
        query=query,
        supplier_id=supplier_id,
    )
    matching_count = base.count()
    if scope == "current_page":
        bounded_size = min(100, max(1, page_size))
        offset = (max(1, page) - 1) * bounded_size
        rows = list(base[offset : offset + bounded_size])
    elif scope == "selected":
        selected = {str(value).strip() for value in worker_ids if str(value).strip()}
        if not selected:
            raise ValidationError({"scope": "Select at least one worker before exporting the selected scope."})
        rows = list(base.filter(pk__in=selected))
    else:
        if matching_count > EXPORT_MAX_ROWS:
            raise ValidationError(
                {"scope": f"This export matches {matching_count:,} workers. Narrow the filters to {EXPORT_MAX_ROWS:,} or fewer workers before exporting."}
            )
        rows = list(base)
    if len(rows) > EXPORT_MAX_ROWS:
        raise ValidationError({"scope": f"An export can contain at most {EXPORT_MAX_ROWS:,} worker rows."})
    return rows, matching_count


def build_rental_timesheet_export_dataset(
    *,
    company,
    membership,
    project_id,
    period_start: date,
    columns: Iterable[str],
    query: str = "",
    supplier_id: str = "",
    scope: str = "all_matching",
    page: int = 1,
    page_size: int = 50,
    worker_ids: Iterable[str] = (),
) -> dict[str, Any]:
    if scope not in VALID_EXPORT_SCOPES:
        raise ValidationError({"scope": "Export scope must be all matching, current page, or selected workers."})
    start, end = month_bounds(period_start)
    project = _resolve_project(company=company, membership=membership, project_id=project_id)
    supplier_id = _normalized_supplier_id(supplier_id)
    available = _available_columns(start, membership)
    by_key = {column.key: column for column in available}
    requested_keys = []
    for raw in columns:
        key = str(raw or "").strip()
        if key and key not in requested_keys:
            requested_keys.append(key)
    if not requested_keys:
        raise ValidationError({"columns": "Select at least one export column."})
    unavailable = [key for key in requested_keys if key not in by_key]
    if unavailable:
        raise PermissionDenied("One or more requested export columns are not available to your access profile.")
    selected_columns = [by_key[key] for key in requested_keys]

    workers, matching_count = _select_worker_queryset(
        company=company,
        project=project,
        start=start,
        end=end,
        query=(query or "").strip(),
        supplier_id=supplier_id,
        scope=scope,
        page=page,
        page_size=page_size,
        worker_ids=worker_ids,
    )
    worker_ids_db = [worker.pk for worker in workers]
    period = (
        RentalTimesheetPeriod.objects.for_company(company)
        .select_related("submitted_by", "approved_by", "locked_by")
        .filter(project=project, period_start=start)
        .first()
    )

    policy = effective_timesheet_policy(company=company, project=project, period=period)

    assignments = list(
        WorkerAssignment.objects.for_company(company)
        .filter(worker_id__in=worker_ids_db, project=project, cancelled_at__isnull=True, effective_from__lte=end)
        .filter(Q(effective_to__isnull=True) | Q(effective_to__gte=start))
        .order_by("worker_id", "effective_from", "created_at")
    ) if worker_ids_db else []
    assignments_by_worker: dict[str, list[WorkerAssignment]] = {}
    for assignment in assignments:
        assignments_by_worker.setdefault(str(assignment.worker_id), []).append(assignment)

    entries = list(
        RentalTimesheetEntry.objects.for_company(company)
        .filter(period=period, worker_id__in=worker_ids_db)
        .select_related("assignment")
        .order_by("worker_id", "work_date")
    ) if period and worker_ids_db else []
    entries_by_worker: dict[str, dict[int, RentalTimesheetEntry]] = {}
    for entry in entries:
        entries_by_worker.setdefault(str(entry.worker_id), {})[entry.work_date.day] = entry

    overtime_rows = list(
        RentalTimesheetOvertime.objects.for_company(company).filter(period=period, worker_id__in=worker_ids_db)
    ) if period and worker_ids_db else []
    overtime_by_worker = {str(row.worker_id): row for row in overtime_rows}

    workflow = {
        "timesheet_status": period.get_status_display() if period else "Draft",
        "timesheet_revision": period.revision if period else 0,
        "submitted_at": _date_time_label(period.submitted_at) if period else "",
        "submitted_by": _actor_label(period.submitted_by) if period else "",
        "approved_at": _date_time_label(period.approved_at) if period else "",
        "approved_by": _actor_label(period.approved_by) if period else "",
        "locked_at": _date_time_label(period.locked_at) if period else "",
        "locked_by": _actor_label(period.locked_by) if period else "",
    }
    project_values = {
        "project_code": project.code,
        "project_name": project.name,
        "project_client": project.client_name,
        "project_location": project.location,
        "project_manager": project.manager_name,
        "period": f"{start:%B %Y}",
    }

    rows: list[dict[str, Any]] = []
    for worker in workers:
        wid = str(worker.pk)
        supplier = worker.supplier
        worker_assignments = assignments_by_worker.get(wid, [])
        worker_entries = entries_by_worker.get(wid, {})
        overtime = overtime_by_worker.get(wid)
        assigned_days = 0
        regular_hours = Decimal("0")
        automatic_overtime_hours = Decimal("0")
        automatic_overtime_wage = Decimal("0")
        base_wage = Decimal("0")
        work_days = absent = no_scope = leave = off = zero = missing = 0
        notes: list[str] = []
        day_values: dict[str, str] = {}
        for day in range(1, end.day + 1):
            work_date = start.replace(day=day)
            assigned = next(
                (
                    assignment
                    for assignment in worker_assignments
                    if assignment.effective_from <= work_date
                    and (assignment.effective_to is None or assignment.effective_to >= work_date)
                ),
                None,
            )
            entry = worker_entries.get(day)
            key = f"day_{day:02d}"
            if assigned is None:
                day_values[key] = "Not assigned"
                continue
            assigned_days += 1
            if entry is None:
                day_values[key] = ""
                missing += 1
                continue
            day_values[key] = _entry_display(entry)
            if entry.note:
                notes.append(f"{day:02d}: {entry.note}")
            if entry.code == "A":
                absent += 1
            elif entry.code == "N":
                no_scope += 1
            elif entry.code == "L":
                leave += 1
            elif entry.code == "OFF":
                off += 1
            else:
                regular_part, automatic_ot_part = split_daily_hours(entry.regular_hours, policy)
                regular_hours += regular_part
                automatic_overtime_hours += automatic_ot_part
                if entry.regular_hours > 0:
                    work_days += 1
                else:
                    zero += 1

            # Commercial board/export preview follows the authoritative project-period
            # hours policy. Daily input remains the worker's total worked hours; the
            # configured regular-day threshold splits base vs automatic OT.
            if entry is not None:
                regular_part, automatic_ot_part = split_daily_hours(entry.regular_hours, policy)
                if assigned.rate_type == "hourly":
                    base_wage += regular_part * assigned.rate
                elif assigned.rate_type == "daily" and entry.regular_hours > 0:
                    base_wage += assigned.rate
                elif assigned.rate_type == "monthly":
                    base_wage += assigned.rate / Decimal(end.day)
                if automatic_ot_part > 0:
                    automatic_overtime_wage += automatic_ot_part * overtime_bill_rate(
                        assignment=assigned, policy=policy, period_start=start
                    )

        trade_values = list(dict.fromkeys(a.trade for a in worker_assignments if a.trade))
        rate_type_values = list(dict.fromkeys(a.get_rate_type_display() for a in worker_assignments))
        rate_values = list(dict.fromkeys(str(a.rate) for a in worker_assignments))
        starts = [a.effective_from for a in worker_assignments]
        ends = [a.effective_to for a in worker_assignments]
        additional_overtime_hours = overtime.hours if overtime else Decimal("0")
        additional_overtime_wage = (overtime.hours * overtime.rate) if overtime else Decimal("0")
        total_overtime_hours = automatic_overtime_hours + additional_overtime_hours
        overtime_wage = automatic_overtime_wage + additional_overtime_wage
        derived_rates = {
            overtime_bill_rate(assignment=a, policy=policy, period_start=start)
            for a in worker_assignments
        }
        displayed_overtime_rate = overtime.rate if overtime else (next(iter(derived_rates)) if len(derived_rates) == 1 else Decimal("0"))
        row: dict[str, Any] = {
            "worker_id": worker.worker_number,
            "worker_name": worker.full_name,
            "national_id": worker.national_id,
            "worker_phone": worker.phone,
            "worker_status": worker.get_status_display(),
            "worker_notes": worker.notes,
            "supplier_code": supplier.code,
            "supplier_name": supplier.name,
            "supplier_contact": supplier.contact_person,
            "supplier_phone": supplier.phone,
            "supplier_email": supplier.email,
            "supplier_cr": supplier.cr_number,
            "supplier_vat": supplier.vat_number,
            "supplier_payment_terms": supplier.payment_terms,
            "supplier_address": supplier.address,
            **project_values,
            "regular_hours_limit": float(policy.regular_hours_per_day),
            "assignment_trade": " → ".join(trade_values),
            "assignment_rate_type": " → ".join(rate_type_values),
            "assignment_rate": " → ".join(rate_values),
            "assignment_start": min(starts).isoformat() if starts else "",
            "assignment_end": (
                "Open" if ends and any(value is None for value in ends) else max(value for value in ends if value is not None).isoformat()
            ) if ends else "",
            "assignment_segments": "; ".join(_segment_label(a) for a in worker_assignments),
            **day_values,
            "assigned_days": assigned_days,
            "regular_hours": float(regular_hours),
            "total_hours": float(regular_hours + total_overtime_hours),
            "work_days": work_days,
            "absent_days": absent,
            "no_scope_days": no_scope,
            "leave_days": leave,
            "off_days": off,
            "zero_hour_days": zero,
            "missing_days": missing,
            "daily_notes": "; ".join(notes),
            "overtime_hours": float(total_overtime_hours),
            "automatic_overtime_hours": float(automatic_overtime_hours),
            "additional_overtime_hours": float(additional_overtime_hours),
            "overtime_multiplier": float(policy.overtime_multiplier),
            "overtime_rate": float(displayed_overtime_rate),
            "base_wage": float(base_wage),
            "overtime_wage": float(overtime_wage),
            "gross_wage": float(base_wage + overtime_wage),
            **workflow,
        }
        rows.append(row)

    scope_labels = {
        "all_matching": "All matching workers",
        "current_page": "Current page",
        "selected": "Selected workers",
    }
    return {
        "company": company,
        "project": project,
        "period": period,
        "period_start": start,
        "period_end": end,
        "columns": selected_columns,
        "rows": rows,
        "matching_count": matching_count,
        "scope": scope,
        "scope_label": scope_labels[scope],
        "query": (query or "").strip(),
        "supplier_id": supplier_id,
        "status": workflow["timesheet_status"],
        "revision": workflow["timesheet_revision"],
        "exported_by": _actor_label(getattr(membership, "user", None)),
    }


def _safe_excel_value(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    if value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


def _safe_filename_part(value: object) -> str:
    text = re.sub(r"[^A-Za-z0-9._-]+", "-", str(value or "").strip()).strip("-._")
    return text or "timesheet"


def export_filename(dataset: dict[str, Any], extension: str) -> str:
    project = dataset["project"]
    period_start = dataset["period_start"]
    return f"rental-timesheet-{_safe_filename_part(project.code).lower()}-{period_start:%Y-%m}.{extension}"


def render_rental_timesheet_xlsx(dataset: dict[str, Any]) -> bytes:
    columns: list[ExportColumn] = dataset["columns"]
    rows: list[dict[str, Any]] = dataset["rows"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Timesheet"
    last_col = max(1, len(columns))
    last_letter = get_column_letter(last_col)

    company = dataset["company"]
    project = dataset["project"]
    period_start = dataset["period_start"]
    generated = timezone.localtime().strftime("%Y-%m-%d %H:%M")
    filter_text = "No additional filters"
    bits = []
    if dataset["query"]:
        bits.append(f'Search: {dataset["query"]}')
    if dataset["supplier_id"]:
        bits.append("Supplier filter applied")
    if bits:
        filter_text = " · ".join(bits)

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    sheet.cell(1, 1, f"{company.name} · Rental Manpower Timesheet Export")
    sheet.cell(1, 1).font = Font(bold=True, size=15)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    sheet.cell(2, 1, f"Project: {project.code} · {project.name}   |   Period: {period_start:%B %Y}   |   Status: {dataset['status']} · Rev {dataset['revision']}")
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
    sheet.cell(3, 1, f"Scope: {dataset['scope_label']} · {len(rows):,} exported worker rows   |   {filter_text}")
    sheet.merge_cells(start_row=4, start_column=1, end_row=4, end_column=last_col)
    actor = dataset.get("exported_by") or "Authorized user"
    sheet.cell(4, 1, f"Generated: {generated}   |   Exported by: {actor}")
    sheet.merge_cells(start_row=5, start_column=1, end_row=5, end_column=last_col)
    sheet.cell(5, 1, "Daily entries: hours (0–24) · A = Absent · N = No Scope · L = Leave · OFF = Off · blank = missing on an assigned day · Not assigned = no project assignment")
    sheet.cell(5, 1).font = Font(size=9, italic=True, color="5F6368")
    sheet.cell(5, 1).alignment = Alignment(vertical="center", wrap_text=True)
    for row_index in (2, 3, 4):
        sheet.cell(row_index, 1).font = Font(size=10)
        sheet.cell(row_index, 1).alignment = Alignment(vertical="center")

    header_row = 6
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(header_row, index, column.label)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="202329")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color="8E949B"))
        sheet.column_dimensions[get_column_letter(index)].width = min(48, max(10, column.width))
    sheet.row_dimensions[header_row].height = 30

    numeric_keys = {column.key for column in columns if column.kind == "number"}
    for row_index, item in enumerate(rows, start=header_row + 1):
        for col_index, column in enumerate(columns, start=1):
            value = _safe_excel_value(item.get(column.key, ""))
            cell = sheet.cell(row_index, col_index, value)
            cell.alignment = Alignment(
                horizontal="right" if column.key in numeric_keys else ("center" if column.group == "Daily entries" else "left"),
                vertical="top",
                wrap_text=column.width >= 26,
            )
            if column.key in numeric_keys and isinstance(value, (int, float, Decimal)):
                cell.number_format = "0.00" if isinstance(value, (float, Decimal)) else "0"
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F7F7F7")
    sheet.freeze_panes = f"A{header_row + 1}"
    if rows:
        sheet.auto_filter.ref = f"A{header_row}:{last_letter}{header_row + len(rows)}"
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = f"{header_row}:{header_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddFooter.center.text = "SESCCO MS · Rental Manpower Timesheet"
    sheet.oddFooter.right.text = "Page &P of &N"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _pdf_chunks(columns: list[ExportColumn], width_budget_mm: float = 265.0) -> list[list[ExportColumn]]:
    if not columns:
        return []
    anchors = [column for column in columns if column.key in {"worker_id", "worker_name"}]
    anchor_keys = {column.key for column in anchors}
    anchor_width = sum(column.pdf_width_mm for column in anchors)
    chunks: list[list[ExportColumn]] = []
    current: list[ExportColumn] = list(anchors)
    used = anchor_width
    for column in columns:
        if column.key in anchor_keys:
            continue
        width = max(12.0, column.pdf_width_mm)
        if len(current) > len(anchors) and used + width > width_budget_mm:
            chunks.append(current)
            current = list(anchors)
            used = anchor_width
        current.append(column)
        used += width
    if current and (len(current) > len(anchors) or not chunks):
        chunks.append(current)
    return chunks


def _pdf_cell(value: Any, style: ParagraphStyle) -> Paragraph:
    text = "" if value is None else str(value)
    return Paragraph(html_escape(text).replace("\n", "<br/>"), style)


def render_rental_timesheet_pdf(dataset: dict[str, Any]) -> bytes:
    columns: list[ExportColumn] = dataset["columns"]
    rows: list[dict[str, Any]] = dataset["rows"]
    output = BytesIO()
    page_size = landscape(A4)
    doc = SimpleDocTemplate(
        output,
        pagesize=page_size,
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=12 * mm,
        title=f"{dataset['project'].code} {dataset['period_start']:%Y-%m} Rental Timesheet",
        author=dataset["company"].name,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ExportTitle", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=13, leading=15, spaceAfter=4)
    meta_style = ParagraphStyle("ExportMeta", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.5, leading=9.5, textColor=colors.HexColor("#4B5057"))
    header_style = ParagraphStyle("ExportHeader", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=6.5, leading=7.5, textColor=colors.white, alignment=TA_CENTER)
    cell_style = ParagraphStyle("ExportCell", parent=styles["BodyText"], fontName="Helvetica", fontSize=6.2, leading=7.5, textColor=colors.HexColor("#202329"), alignment=TA_LEFT)
    center_style = ParagraphStyle("ExportCellCenter", parent=cell_style, alignment=TA_CENTER)
    section_style = ParagraphStyle("ExportSection", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=8.5, leading=10, spaceBefore=6, spaceAfter=4)

    project = dataset["project"]
    period_start = dataset["period_start"]
    generated = timezone.localtime().strftime("%Y-%m-%d %H:%M")
    filter_bits = []
    if dataset["query"]:
        filter_bits.append(f"Search: {dataset['query']}")
    if dataset["supplier_id"]:
        filter_bits.append("Supplier filter applied")
    filter_label = " · ".join(filter_bits) if filter_bits else "No additional filters"

    story = [
        Paragraph(f"{html_escape(dataset['company'].name)} · Rental Manpower Timesheet Export", title_style),
        Paragraph(
            f"Project: <b>{html_escape(project.code)} · {html_escape(project.name)}</b> &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"Period: <b>{period_start:%B %Y}</b> &nbsp;&nbsp;|&nbsp;&nbsp; Status: <b>{html_escape(dataset['status'])}</b> · Rev {dataset['revision']}",
            meta_style,
        ),
        Paragraph(
            f"Scope: <b>{html_escape(dataset['scope_label'])}</b> · {len(rows):,} exported worker rows &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"{html_escape(filter_label)} &nbsp;&nbsp;|&nbsp;&nbsp; Generated: {generated} &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"Exported by: {html_escape(dataset.get('exported_by') or 'Authorized user')}",
            meta_style,
        ),
        Paragraph(
            "Daily entries: hours (0–24) · A = Absent · N = No Scope · L = Leave · OFF = Off · "
            "blank = missing on an assigned day · Not assigned = no project assignment",
            meta_style,
        ),
        Spacer(1, 5 * mm),
    ]

    chunks = _pdf_chunks(columns)
    for chunk_index, chunk in enumerate(chunks, start=1):
        if len(chunks) > 1:
            story.append(Paragraph(f"Column section {chunk_index} of {len(chunks)}", section_style))
        table_data = [[_pdf_cell(column.label, header_style) for column in chunk]]
        for item in rows:
            table_data.append([
                _pdf_cell(item.get(column.key, ""), center_style if column.group == "Daily entries" or column.kind == "number" else cell_style)
                for column in chunk
            ])
        widths = [column.pdf_width_mm * mm for column in chunk]
        table = Table(table_data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#202329")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D6D9DD")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F7F7")]),
                    ("LEFTPADDING", (0, 0), (-1, -1), 2.2),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 2.2),
                    ("TOPPADDING", (0, 0), (-1, -1), 2.2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2),
                ]
            )
        )
        story.append(table)
        if chunk_index != len(chunks):
            story.append(Spacer(1, 7 * mm))

    def draw_footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#6B7077"))
        canvas.drawString(10 * mm, 6 * mm, "SESCCO MS · Rental Manpower Timesheet")
        canvas.drawRightString(page_size[0] - 10 * mm, 6 * mm, f"Page {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)
    return output.getvalue()


def render_rental_timesheet_export(dataset: dict[str, Any], export_format: str) -> tuple[bytes, str, str]:
    normalized = str(export_format or "").strip().lower()
    if normalized not in VALID_EXPORT_FORMATS:
        raise ValidationError({"format": "Export format must be xlsx or pdf."})
    if normalized == "xlsx":
        return (
            render_rental_timesheet_xlsx(dataset),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            export_filename(dataset, "xlsx"),
        )
    return render_rental_timesheet_pdf(dataset), "application/pdf", export_filename(dataset, "pdf")
