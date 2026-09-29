import io
import json
from datetime import date

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.models import CompanyMembership, MembershipProjectScope, User
from apps.accounts.roles import AccessRole, ScopeMode
from apps.core.models import AuditEvent, Company
from apps.rental_manpower.services.assignments import assign_worker
from apps.rental_manpower.services.masters import create_project, create_supplier, create_worker
from apps.rental_manpower.services.timesheets import save_entries, save_overtime


class RentalTimesheetExportTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(name="Export Co", slug="rental-export")
        self.user = User.objects.create_user(username="export-owner", password="test-password")
        self.owner = CompanyMembership.objects.create(
            company=self.company,
            user=self.user,
            role=AccessRole.OWNER,
        )
        self.client.force_login(self.user)
        self.supplier = create_supplier(
            actor_membership=self.owner,
            code="SUP-EX",
            name="Export Supplier",
            email="supplier@example.test",
            vat_number="VAT-001",
        )
        self.project = create_project(
            actor_membership=self.owner,
            code="PRJ-EX",
            name="Export Project",
            start_date=date(2026, 8, 1),
        )
        self.worker = create_worker(
            actor_membership=self.owner,
            supplier_id=self.supplier.pk,
            worker_number="RW-EX-001",
            full_name="Export Worker",
            national_id="NID-EX-001",
            phone="0500000001",
        )
        self.worker_two = create_worker(
            actor_membership=self.owner,
            supplier_id=self.supplier.pk,
            worker_number="RW-EX-002",
            full_name="Second Worker",
        )
        for worker in (self.worker, self.worker_two):
            assign_worker(
                actor_membership=self.owner,
                worker_id=worker.pk,
                project_id=self.project.pk,
                trade="Driver",
                rate_type="Hourly",
                rate="10",
                effective_date=date(2026, 8, 1),
            )
        save_entries(
            actor_membership=self.owner,
            project_id=self.project.pk,
            period_start=date(2026, 8, 1),
            entries=[
                {"worker_id": self.worker.pk, "work_date": date(2026, 8, 1), "value": "8"},
                {"worker_id": self.worker.pk, "work_date": date(2026, 8, 2), "value": "A"},
            ],
        )
        save_overtime(
            actor_membership=self.owner,
            project_id=self.project.pk,
            period_start=date(2026, 8, 1),
            worker_id=self.worker.pk,
            hours="2",
        )
        self.url = reverse("rental_manpower:timesheets-export-api")

    def test_export_schema_is_dynamic_and_includes_month_days(self):
        response = self.client.get(
            self.url,
            {"project_id": str(self.project.reference), "period": "2026-08"},
        )
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["matchingCount"], 2)
        self.assertEqual([item["value"] for item in payload["formats"]], ["xlsx", "pdf"])
        keys = [column["key"] for group in payload["groups"] for column in group["columns"]]
        self.assertIn("day_01", keys)
        self.assertIn("day_31", keys)
        self.assertIn("supplier_vat", keys)
        self.assertIn("assignment_rate", keys)
        self.assertIn("total_hours", keys)
        self.assertIn("base_wage", keys)
        self.assertIn("overtime_wage", keys)
        self.assertIn("gross_wage", keys)
        self.assertIn("day_01", payload["defaultColumns"])
        self.assertIn("regular_hours", payload["defaultColumns"])

    def test_xlsx_export_contains_all_matching_rows_and_selected_columns(self):
        response = self.client.post(
            self.url,
            data=json.dumps(
                {
                    "project_id": str(self.project.reference),
                    "period": "2026-08",
                    "format": "xlsx",
                    "scope": "all_matching",
                    "columns": [
                        "worker_id",
                        "worker_name",
                        "supplier_name",
                        "day_01",
                        "day_02",
                        "regular_hours",
                        "overtime_hours",
                        "missing_days",
                    ],
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertEqual(response["X-Export-Row-Count"], "2")
        workbook = load_workbook(io.BytesIO(response.content), data_only=True)
        sheet = workbook["Timesheet"]
        self.assertIn("A = Absent", sheet.cell(5, 1).value)
        self.assertIn("Exported by: export-owner", sheet.cell(4, 1).value)
        headers = [sheet.cell(6, index).value for index in range(1, 9)]
        self.assertEqual(
            headers,
            ["Worker ID", "Worker name", "Supplier name", "01 Sat", "02 Sun", "Regular hours", "OT hours", "Missing days"],
        )
        exported = {sheet.cell(row, 1).value: [sheet.cell(row, col).value for col in range(1, 9)] for row in range(7, 9)}
        self.assertEqual(exported["RW-EX-001"][3], "8")
        self.assertEqual(exported["RW-EX-001"][4], "A")
        self.assertEqual(exported["RW-EX-001"][5], 8)
        self.assertEqual(exported["RW-EX-001"][6], 2)
        self.assertTrue(AuditEvent.objects.filter(company=self.company, action="rental.timesheet.exported").exists())

    def test_pdf_export_is_a_real_pdf_file(self):
        response = self.client.post(
            self.url,
            data=json.dumps(
                {
                    "project_id": str(self.project.reference),
                    "period": "2026-08",
                    "format": "pdf",
                    "scope": "current_page",
                    "page": 1,
                    "page_size": 25,
                    "columns": ["worker_id", "worker_name", "supplier_name", "day_01", "regular_hours"],
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF-"))
        self.assertIn(".pdf", response["Content-Disposition"])

    def test_selected_scope_only_exports_requested_matching_workers(self):
        response = self.client.post(
            self.url,
            data=json.dumps(
                {
                    "project_id": str(self.project.reference),
                    "period": "2026-08",
                    "format": "xlsx",
                    "scope": "selected",
                    "worker_ids": [str(self.worker_two.pk)],
                    "columns": ["worker_id", "worker_name"],
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["X-Export-Row-Count"], "1")
        workbook = load_workbook(io.BytesIO(response.content), data_only=True)
        sheet = workbook["Timesheet"]
        self.assertEqual(sheet.cell(7, 1).value, "RW-EX-002")
        self.assertEqual(sheet.cell(8, 1).value, None)

    def test_supervisor_schema_does_not_offer_supplier_private_or_commercial_columns(self):
        supervisor_user = User.objects.create_user(username="export-supervisor", password="test-password")
        supervisor = CompanyMembership.objects.create(
            company=self.company,
            user=supervisor_user,
            role=AccessRole.RENTAL_SUPERVISOR,
            project_scope_mode=ScopeMode.SELECTED,
        )
        MembershipProjectScope.objects.create(membership=supervisor, project=self.project)
        self.client.force_login(supervisor_user)
        response = self.client.get(
            self.url,
            {"project_id": str(self.project.reference), "period": "2026-08"},
        )
        self.assertEqual(response.status_code, 200)
        keys = [column["key"] for group in response.json()["groups"] for column in group["columns"]]
        self.assertNotIn("supplier_email", keys)
        self.assertNotIn("supplier_vat", keys)
        self.assertNotIn("assignment_rate", keys)
        self.assertIn("worker_phone", keys)

        blocked = self.client.post(
            self.url,
            data=json.dumps(
                {
                    "project_id": str(self.project.reference),
                    "period": "2026-08",
                    "format": "xlsx",
                    "scope": "all_matching",
                    "columns": ["worker_id", "assignment_rate"],
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(blocked.status_code, 403)
