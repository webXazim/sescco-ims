from __future__ import annotations

import io

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from apps.sourcing.tests.support import SOURCING_HTTP_TEST_STORAGES
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.access_catalog import AccessPermission
from apps.accounts.models import AccessProfile, AccessProfilePermission, CompanyMembership
from apps.accounts.roles import AccessRole
from apps.core.models import AuditArea, AuditEvent, Company
from apps.sourcing.exchange import HEADERS, import_sourcing_rows
from apps.sourcing.models import (
    SourcingMaterial,
    SourcingManpowerSupplier,
    SourcingTrade,
    SourcingVendor,
    SourcingVendorOffer,
    SourcingVendorOfferRevision,
    SourcingWorkforceOffer,
)


@override_settings(SECURE_SSL_REDIRECT=False, STORAGES=SOURCING_HTTP_TEST_STORAGES)
class SourcingDataExchangeTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.company = Company.objects.create(name="Sourcing Exchange Co", slug="sourcing-exchange-co")
        self.owner_user = User.objects.create_user(username="exchange-owner", password="strong-test-password")
        self.owner = CompanyMembership.objects.create(company=self.company, user=self.owner_user, role=AccessRole.OWNER)

    def _member(self, username: str, permissions: tuple[str, ...]):
        User = get_user_model()
        user = User.objects.create_user(username=username, password="strong-test-password")
        profile = AccessProfile.objects.create(company=self.company, key=f"p-{username}", name=f"{username} profile", is_system=False, is_active=True)
        AccessProfilePermission.objects.bulk_create(AccessProfilePermission(profile=profile, permission=p) for p in permissions)
        membership = CompanyMembership.objects.create(company=self.company, user=user, role=AccessRole.CUSTOM, access_profile=profile)
        return user, membership

    def _upload(self, text: str, name="import.csv"):
        return SimpleUploadedFile(name, text.encode("utf-8"), content_type="text/csv")

    def test_vendor_import_dry_run_rolls_back_and_commit_is_audited(self):
        self.client.force_login(self.owner_user)
        body = (
            "code,name,display_name,mobile,email,address,city,region,cr_number,vat_number,status\n"
            "VND-100,Test Vendor,Test Vendor,+966500000100,test.vendor@example.com,King Fahd Road,Dammam,Eastern Province,CR-100,VAT-100,active\n"
        )
        response = self.client.post(reverse("sourcing:data_exchange"), {"dataset": "vendors", "dry_run": "on", "file": self._upload(body)})
        self.assertEqual(response.status_code, 200)
        self.assertFalse(SourcingVendor.objects.filter(company=self.company, code="VND-100").exists())
        self.assertContains(response, "Validation passed")

        response = self.client.post(reverse("sourcing:data_exchange"), {"dataset": "vendors", "file": self._upload(body)})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(SourcingVendor.objects.filter(company=self.company, code="VND-100", city="Dammam").exists())
        self.assertTrue(AuditEvent.objects.filter(company=self.company, area=AuditArea.SOURCING, action="sourcing.data_import.completed").exists())

    def test_duplicate_import_identity_rejects_whole_batch(self):
        headers = ["code", "name"]
        rows = [{"code": "MAT-1", "name": "Rebar"}, {"code": "MAT-1", "name": "Different"}]
        result = import_sourcing_rows(actor_membership=self.owner, dataset="materials", headers=headers, records=rows, dry_run=False)
        self.assertEqual(len(result.errors), 1)
        self.assertFalse(SourcingMaterial.objects.filter(company=self.company).exists())

    def test_catalog_import_can_bulk_verify_reference_and_preserves_revision(self):
        vendor = SourcingVendor.objects.create(company=self.company, code="VND-A", name="Vendor A")
        material = SourcingMaterial.objects.create(company=self.company, code="MAT-A", name="Rebar", default_unit="ton")
        headers = ["vendor_code", "material_code", "availability", "available_quantity", "unit", "rate", "currency", "verified_now", "contact_name", "verification_note"]
        rows = [{"vendor_code": vendor.code, "material_code": material.code, "availability": "available", "available_quantity": "25", "unit": "ton", "rate": "2180", "currency": "SAR", "verified_now": "yes", "contact_name": "Mohammed", "verification_note": "Confirmed by phone"}]
        result = import_sourcing_rows(actor_membership=self.owner, dataset="vendor_catalog", headers=headers, records=rows, dry_run=False)
        self.assertFalse(result.errors)
        offer = SourcingVendorOffer.objects.get(company=self.company, vendor=vendor, material=material)
        self.assertEqual(str(offer.rate), "2180.0000")
        self.assertIsNotNone(offer.last_verified_at)
        revision = SourcingVendorOfferRevision.objects.get(offer_id_snapshot=offer.pk)
        self.assertEqual(revision.contact_name, "Mohammed")

    def test_workforce_import_is_sourcing_only(self):
        supplier = SourcingManpowerSupplier.objects.create(company=self.company, code="MPS-A", name="Supplier A")
        trade = SourcingTrade.objects.create(company=self.company, code="TRD-A", name="AC Technician")
        headers = ["supplier_code", "trade_code", "availability", "available_quantity", "rate", "currency", "rate_basis", "verified_now"]
        rows = [{"supplier_code": supplier.code, "trade_code": trade.code, "availability": "available", "available_quantity": "8", "rate": "3200", "currency": "SAR", "rate_basis": "month", "verified_now": "yes"}]
        result = import_sourcing_rows(actor_membership=self.owner, dataset="workforce_catalog", headers=headers, records=rows, dry_run=False)
        self.assertFalse(result.errors)
        self.assertEqual(SourcingWorkforceOffer.objects.filter(company=self.company).count(), 1)
        # The Sourcing app carries no Rental Worker model/relation; this test intentionally stops at Sourcing-owned rows.

    def test_export_requires_export_permission_plus_dataset_view(self):
        viewer, _ = self._member("vendor-viewer-no-export", (AccessPermission.SOURCING_VENDORS_VIEW.value,))
        SourcingVendor.objects.create(company=self.company, code="VND-X", name="Vendor X")
        self.client.force_login(viewer)
        self.assertEqual(self.client.get(reverse("sourcing:data_export", args=["vendors"]), {"format": "csv"}).status_code, 403)

        exporter, _ = self._member("vendor-exporter", (AccessPermission.SOURCING_VENDORS_VIEW.value, AccessPermission.SOURCING_EXPORT_EXECUTE.value))
        self.client.force_login(exporter)
        response = self.client.get(reverse("sourcing:data_export", args=["vendors"]), {"format": "csv"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["X-Content-Type-Options"], "nosniff")
        self.assertIn("VND-X", response.content.decode("utf-8-sig"))

    def test_xlsx_export_is_readable_and_formula_cells_are_neutralized(self):
        SourcingVendor.objects.create(company=self.company, code="VND-F", name="Formula Vendor", notes="=HYPERLINK(\"https://example.test\")")
        self.client.force_login(self.owner_user)
        response = self.client.get(reverse("sourcing:data_export", args=["vendors"]), {"format": "xlsx"})
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content), read_only=True, data_only=False)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        workbook.close()
        notes_index = HEADERS["vendors"].index("notes")
        self.assertTrue(str(values[1][notes_index]).startswith("'="))
