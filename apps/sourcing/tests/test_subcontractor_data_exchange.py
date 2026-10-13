"""Upgrade 08: master/service catalog bulk import, export, and permission tests."""
from __future__ import annotations

import csv
import io

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.access_catalog import AccessPermission
from apps.accounts.models import AccessProfile, AccessProfilePermission, CompanyMembership
from apps.accounts.roles import AccessRole
from apps.core.models import AuditArea, AuditEvent, Company
from apps.sourcing.exchange import HEADERS, import_sourcing_rows
from apps.sourcing.models import (
    SourcingService, SourcingSubcontractor, SourcingSubcontractorService,
    SourcingSubcontractorServiceRevision,
)
from apps.sourcing.tests.support import SOURCING_HTTP_TEST_STORAGES


@override_settings(SECURE_SSL_REDIRECT=False, STORAGES=SOURCING_HTTP_TEST_STORAGES)
class SubcontractorExchangeTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(name="Subcontractor Exchange One", slug="sub-exchange-one")
        self.other = Company.objects.create(name="Subcontractor Exchange Two", slug="sub-exchange-two")
        self.user = get_user_model().objects.create_user(username="sub-exchange-owner", password="strong-test-password")
        self.member = CompanyMembership.objects.create(company=self.company, user=self.user, role=AccessRole.OWNER)
        self.client.force_login(self.user)

    def _member(self, username, permissions):
        user = get_user_model().objects.create_user(username=username, password="strong-test-password")
        profile = AccessProfile.objects.create(company=self.company, key=username, name=username, is_active=True)
        AccessProfilePermission.objects.bulk_create([
            AccessProfilePermission(profile=profile, permission=perm.value) for perm in permissions
        ])
        CompanyMembership.objects.create(company=self.company, user=user, role=AccessRole.CUSTOM,
                                         access_profile=profile)
        return user

    def _csv_upload(self, data):
        return SimpleUploadedFile("subcontractor-import.csv", data.encode("utf-8"), content_type="text/csv")

    def test_master_dry_run_then_commit_and_alias_roundtrip(self):
        services = [{"code": "SRV-HVAC", "name": "Air Conditioning Installation", "category": "HVAC",
                     "aliases": "AC Installation | Cooling System", "default_rate_basis": "job"}]
        result = import_sourcing_rows(actor_membership=self.member, dataset="services",
                                      headers=HEADERS["services"], records=services, dry_run=True)
        self.assertFalse(result.errors)
        self.assertEqual(result.creates, 1)
        self.assertFalse(SourcingService.objects.filter(company=self.company).exists())
        result = import_sourcing_rows(actor_membership=self.member, dataset="services",
                                      headers=HEADERS["services"], records=services, dry_run=False)
        self.assertFalse(result.errors)
        self.assertEqual(SourcingService.objects.get(company=self.company).aliases,
                         ["AC Installation", "Cooling System"])
        rows = [{"code": "SUB-100", "name": "Building Experts", "city": "Dammam", "service_coverage": "Eastern Province"}]
        result = import_sourcing_rows(actor_membership=self.member, dataset="subcontractors",
                                      headers=HEADERS["subcontractors"], records=rows, dry_run=False)
        self.assertFalse(result.errors)
        obj = SourcingSubcontractor.objects.get(company=self.company, code="SUB-100")
        self.assertEqual(obj.service_coverage, "Eastern Province")
        self.assertIsNone(obj.last_verified_at)
        self.assertTrue(AuditEvent.objects.filter(company=self.company, area=AuditArea.SOURCING,
                           action="sourcing.data_import.completed", metadata__dataset="subcontractors").exists())

    def test_catalog_multi_scope_verification_and_read_only_timestamp(self):
        sub = SourcingSubcontractor.objects.create(company=self.company, code="SUB-A", name="Contractor A")
        service = SourcingService.objects.create(company=self.company, code="SRV-A", name="HVAC Service")
        rows = [
            {"subcontractor_code": sub.code, "service_code": service.code,
             "scope_description": "Chiller overhaul", "service_location": "Dammam",
             "rate": "300", "rate_basis": "day", "verified_now": "yes", "contact_name": "Site Lead"},
            {"subcontractor_code": sub.code, "service_code": service.code,
             "scope_description": "Air duct installation", "service_location": "Khobar",
             "rate": "500", "rate_basis": "job"},
        ]
        result = import_sourcing_rows(actor_membership=self.member, dataset="subcontractor_service_catalog",
                                      headers=HEADERS["subcontractor_service_catalog"], records=rows, dry_run=False)
        self.assertFalse(result.errors)
        self.assertEqual(result.creates, 2)
        self.assertEqual(SourcingSubcontractorService.objects.filter(company=self.company).count(), 2)
        verified = SourcingSubcontractorService.objects.get(company=self.company, scope_description="Chiller overhaul")
        self.assertIsNotNone(verified.last_verified_at)
        self.assertIsNone(sub.last_verified_at)
        self.assertEqual(SourcingSubcontractorServiceRevision.objects.get(offer=verified).contact_name, "Site Lead")
        update = dict(rows[0], verified_now="no", rate="320", last_verified_at="2035-01-01")
        result = import_sourcing_rows(actor_membership=self.member, dataset="subcontractor_service_catalog",
                                      headers=HEADERS["subcontractor_service_catalog"], records=[update], dry_run=False)
        self.assertFalse(result.errors)
        self.assertEqual(result.updates, 1)
        verified.refresh_from_db()
        self.assertIsNone(verified.last_verified_at)
        self.assertEqual(SourcingSubcontractorServiceRevision.objects.filter(offer=verified).count(), 1)

    def test_duplicate_catalog_identity_causes_transaction_rollback(self):
        sub = SourcingSubcontractor.objects.create(company=self.company, code="SUB-1", name="Company One")
        service = SourcingService.objects.create(company=self.company, code="SRV-1", name="Painting")
        rows = [
            {"subcontractor_code": sub.code, "service_code": service.code,
             "scope_description": "Walls", "service_location": "Dammam"},
            {"subcontractor_code": sub.code.lower(), "service_code": service.code.lower(),
             "scope_description": "WALLS", "service_location": "dammam"},
        ]
        result = import_sourcing_rows(actor_membership=self.member, dataset="subcontractor_service_catalog",
                                      headers=HEADERS["subcontractor_service_catalog"], records=rows, dry_run=False)
        self.assertEqual(len(result.errors), 1)
        self.assertFalse(SourcingSubcontractorService.objects.filter(company=self.company).exists())

    def test_malformed_rate_rejects_batch_without_committing_any_entry(self):
        sub = SourcingSubcontractor.objects.create(company=self.company, code="SUB-20", name="Rate Test")
        service = SourcingService.objects.create(company=self.company, code="SRV-20", name="Rate Service")
        rows = [
            {"subcontractor_code": sub.code, "service_code": service.code,
             "scope_description": "Good", "service_location": "Dammam", "rate": "120"},
            {"subcontractor_code": sub.code, "service_code": service.code,
             "scope_description": "Bad", "service_location": "Dammam", "rate": "1e99999"},
        ]
        result = import_sourcing_rows(actor_membership=self.member, dataset="subcontractor_service_catalog",
                                      headers=HEADERS["subcontractor_service_catalog"], records=rows, dry_run=False)
        self.assertEqual(len(result.errors), 1)
        self.assertFalse(SourcingSubcontractorService.objects.filter(company=self.company).exists())

    def test_cross_company_service_code_is_not_resolved(self):
        sub = SourcingSubcontractor.objects.create(company=self.company, code="SUB-1", name="Company One")
        SourcingService.objects.create(company=self.other, code="SRV-REMOTE", name="Remote Maintenance")
        result = import_sourcing_rows(actor_membership=self.member, dataset="subcontractor_service_catalog",
                                      headers=HEADERS["subcontractor_service_catalog"], records=[
            {"subcontractor_code": sub.code, "service_code": "SRV-REMOTE",
             "scope_description": "Repair", "service_location": "Dammam"}
        ], dry_run=False)
        self.assertEqual(len(result.errors), 1)
        self.assertFalse(SourcingSubcontractorService.objects.filter(company=self.company).exists())

    def test_view_only_export_authorization_and_templates(self):
        viewer = self._member("sub-view-export", (
            AccessPermission.SOURCING_SUBCONTRACTORS_VIEW, AccessPermission.SOURCING_EXPORT_EXECUTE,
        ))
        self.client.force_login(viewer)
        response = self.client.get(reverse("sourcing:data_export", args=["subcontractors"]), {"format": "csv"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.client.get(reverse("sourcing:data_export", args=["services"])).status_code, 403)
        self.assertEqual(self.client.get(reverse("sourcing:data_import_template", args=["subcontractors"])).status_code, 403)
        self.assertEqual(self.client.post(reverse("sourcing:data_exchange"), {
            "dataset": "subcontractors", "file": self._csv_upload("code,name\nSUB-NEW,Not Allowed\n")
        }).status_code, 200)
        self.assertFalse(SourcingSubcontractor.objects.filter(company=self.company, code="SUB-NEW").exists())

        manager = self._member("sub-manage-no-export", (AccessPermission.SOURCING_SUBCONTRACTORS_VIEW,
                                                         AccessPermission.SOURCING_SUBCONTRACTORS_MANAGE))
        self.client.force_login(manager)
        self.assertEqual(self.client.get(reverse("sourcing:data_export", args=["subcontractors"])).status_code, 403)
        template = self.client.get(reverse("sourcing:data_import_template", args=["subcontractors"]), {"format": "csv"})
        self.assertEqual(template.status_code, 200)
        self.assertEqual(template["X-Content-Type-Options"], "nosniff")
        self.assertEqual(next(csv.reader(io.StringIO(template.content.decode("utf-8-sig")))),
                         HEADERS["subcontractors"])
        self.assertEqual(self.client.get(reverse("sourcing:data_import_template", args=["services"])).status_code, 403)

    def test_filtered_service_catalog_export_and_spreadsheet_escaping(self):
        sub = SourcingSubcontractor.objects.create(company=self.company, code="SUB-9", name="Local Contractor")
        remote_sub = SourcingSubcontractor.objects.create(company=self.other, code="SUB-9", name="Remote Contractor")
        svc = SourcingService.objects.create(company=self.company, code="SRV-9", name="Mechanical Overhaul")
        remote_svc = SourcingService.objects.create(company=self.other, code="SRV-9", name="Remote Service")
        SourcingSubcontractorService.objects.create(company=self.company, subcontractor=sub, service=svc,
            scope_description="=HYPERLINK(\"https://example.test\")", service_location="Dammam", rate=300,
            rate_basis="day", is_active=True)
        SourcingSubcontractorService.objects.create(company=self.other, subcontractor=remote_sub, service=remote_svc,
            scope_description="Private", service_location="Dammam", rate=100)
        response = self.client.get(reverse("sourcing:data_export", args=["subcontractor_service_catalog"]),
                                   {"format": "xlsx", "location": "Dammam"})
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content), read_only=True, data_only=False)
        data = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
        self.assertEqual(len(data), 2)
        self.assertEqual(data[1][HEADERS["subcontractor_service_catalog"].index("subcontractor_code")], "SUB-9")
        self.assertTrue(data[1][HEADERS["subcontractor_service_catalog"].index("scope_description")].startswith("'="))

    def test_new_dataset_import_template_xlsx(self):
        response = self.client.get(reverse("sourcing:data_import_template", args=["subcontractor_service_catalog"]))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content), read_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
        self.assertEqual(len(rows), 1)
        self.assertEqual(list(rows[0]), HEADERS["subcontractor_service_catalog"])
