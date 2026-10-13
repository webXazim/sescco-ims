from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.db.models import Exists, OuterRef, Q
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from apps.core.models import AuditArea
from apps.core.services.audit import record_audit_event

from .freshness import (
    FRESHNESS_FRESH, FRESHNESS_NEEDS_VERIFICATION, FRESHNESS_NEVER, FRESHNESS_STALE,
    sourcing_freshness_policy,
)
from .models import (
    SourcingAvailability,
    SourcingEntityStatus,
    SourcingMaterial,
    SourcingService, SourcingServiceRateBasis,
    SourcingSubcontractor, SourcingSubcontractorContact, SourcingSubcontractorService,
    SourcingManpowerSupplier, SourcingManpowerContact,
    SourcingRateBasis,
    SourcingTrade,
    SourcingVendor, SourcingVendorContact,
    SourcingVendorOffer,
    SourcingWorkforceOffer,
)
from .models.base import normalize_text
from .services.catalog import (
    create_material,
    create_vendor_offer,
    set_material_active,
    set_vendor_offer_active,
    update_material,
    update_vendor_offer,
)
from .services.manpower import (
    create_manpower_supplier,
    set_manpower_supplier_status,
    update_manpower_supplier,
)
from .services.vendors import create_vendor, set_vendor_status, update_vendor
from .services.services_master import create_service, update_service, set_service_active
from .services.subcontractors import create_subcontractor, update_subcontractor, set_subcontractor_status
from .services.subcontractor_catalog import (
    create_service_offer, update_service_offer, set_service_offer_active,
)
from .services.workforce import (
    create_trade,
    create_workforce_offer,
    set_trade_active,
    set_workforce_offer_active,
    update_trade,
    update_workforce_offer,
)

MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_IMPORT_ROWS = 5000
IMPORT_DATASETS = (
    ("vendors", "Vendor master"),
    ("materials", "Material master"),
    ("vendor_catalog", "Vendor supply catalog"),
    ("manpower_suppliers", "Manpower supplier master"),
    ("trades", "Worker type / trade master"),
    ("workforce_catalog", "Workforce catalog"),
    ("subcontractors", "Subcontractor master"),
    ("services", "Service master"),
    ("subcontractor_service_catalog", "Subcontractor Service Catalog"),
)
EXPORT_DATASETS = IMPORT_DATASETS

HEADERS = {
    "vendors": ["code", "name", "display_name", "cr_number", "vat_number", "company_phone", "company_email", "website", "primary_contact_name", "mobile", "email", "address", "street_number", "district", "city", "region", "postal_code", "status", "notes", "last_verified_at"],
    "materials": ["code", "name", "category", "default_unit", "aliases", "is_active", "notes"],
    "vendor_catalog": ["vendor_code", "material_code", "specification", "brand", "model", "availability", "available_quantity", "unit", "minimum_quantity", "rate", "currency", "rate_valid_until", "lead_time", "verified_now", "contact_name", "verification_note", "is_active", "notes", "last_verified_at"],
    "manpower_suppliers": ["code", "name", "primary_contact_name", "phone", "mobile", "email", "city", "region", "cr_number", "vat_number", "status", "notes", "last_verified_at"],
    "trades": ["code", "name", "category", "aliases", "is_active", "notes"],
    "workforce_catalog": ["supplier_code", "trade_code", "availability", "available_quantity", "rate", "currency", "rate_basis", "overtime_rate", "rate_valid_until", "mobilization_lead_time", "work_location", "verified_now", "contact_name", "verification_note", "is_active", "notes", "last_verified_at"],
    "subcontractors": ["code", "name", "display_name", "primary_contact_name", "company_phone", "company_email", "mobile", "email", "website", "address", "street_number", "district", "city", "region", "postal_code", "service_coverage", "cr_number", "vat_number", "status", "notes", "last_verified_at"],
    "services": ["code", "name", "category", "default_rate_basis", "aliases", "description", "is_active", "notes"],
    "subcontractor_service_catalog": ["subcontractor_code", "service_code", "scope_description", "service_location", "availability", "rate", "currency", "rate_basis", "rate_valid_until", "lead_time", "verified_now", "contact_name", "verification_note", "is_active", "notes", "last_verified_at"],
}

REQUIRED_HEADERS = {
    "vendors": {"code", "name", "display_name", "mobile", "email", "address", "cr_number", "vat_number"},
    "materials": {"code", "name"},
    "vendor_catalog": {"vendor_code", "material_code"},
    "manpower_suppliers": {"code", "name"},
    "trades": {"code", "name"},
    "workforce_catalog": {"supplier_code", "trade_code"},
    "subcontractors": {"code", "name"},
    "services": {"code", "name"},
    "subcontractor_service_catalog": {"subcontractor_code", "service_code", "scope_description", "service_location"},
}

_HEADER_RE = re.compile(r"[^a-z0-9]+")
_FORMULA_PREFIXES = ("=", "+", "-", "@")


@dataclass(frozen=True)
class ImportRowResult:
    row_number: int
    outcome: str
    identity: str
    message: str = ""


@dataclass(frozen=True)
class ImportResult:
    dataset: str
    dry_run: bool
    rows: tuple[ImportRowResult, ...]

    @property
    def errors(self) -> tuple[ImportRowResult, ...]:
        return tuple(row for row in self.rows if row.outcome == "error")

    @property
    def creates(self) -> int:
        return sum(row.outcome == "create" for row in self.rows)

    @property
    def updates(self) -> int:
        return sum(row.outcome == "update" for row in self.rows)

    @property
    def total(self) -> int:
        return len(self.rows)


class SourcingImportError(ValidationError):
    pass


def _header(value: object) -> str:
    return _HEADER_RE.sub("_", str(value or "").strip().casefold()).strip("_")


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).strip().split())


def _bool(value: object, *, default=False) -> bool:
    raw = _text(value).casefold()
    if raw == "":
        return bool(default)
    if raw in {"1", "true", "yes", "y", "active", "on"}:
        return True
    if raw in {"0", "false", "no", "n", "inactive", "off"}:
        return False
    raise ValidationError(f"Expected yes/no value, got {value!r}.")


def _decimal(value: object, *, field: str) -> Decimal | None:
    raw = _text(value).replace(",", "")
    if not raw:
        return None
    try:
        parsed = Decimal(raw)
    except InvalidOperation as exc:
        raise ValidationError({field: f"Invalid decimal value {value!r}."}) from exc
    if not parsed.is_finite():
        raise ValidationError({field: "A finite decimal is required."})
    if parsed < 0:
        raise ValidationError({field: "Value cannot be negative."})
    return parsed



def _service_rate(value: object) -> Decimal | None:
    """Validate the Service rate against its DecimalField(18,4) before model save."""
    rate = _decimal(value, field="rate")
    if rate is not None and (
        rate >= Decimal("100000000000000") or rate.as_tuple().exponent < -4
    ):
        raise ValidationError({"rate": "Use at most 14 integer digits and 4 decimal places."})
    return rate

def _integer(value: object, *, field: str) -> int | None:
    raw = _text(value).replace(",", "")
    if not raw:
        return None
    try:
        parsed = int(Decimal(raw))
    except (InvalidOperation, ValueError) as exc:
        raise ValidationError({field: f"Invalid whole-number value {value!r}."}) from exc
    if parsed < 0:
        raise ValidationError({field: "Value cannot be negative."})
    return parsed


def _date(value: object, *, field: str) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = _text(value)
    try:
        return date.fromisoformat(raw)
    except ValueError as exc:
        raise ValidationError({field: "Use YYYY-MM-DD."}) from exc


def _aliases(value: object) -> list[str]:
    raw = str(value or "")
    result = []
    seen = set()
    for part in raw.replace("\n", "|").replace(",", "|").split("|"):
        item = _text(part)
        key = item.casefold()
        if item and key not in seen:
            result.append(item)
            seen.add(key)
    return result


def _choice(value: object, *, field: str, choices, default: str) -> str:
    raw = _text(value).casefold() or default
    allowed = {str(key): str(label) for key, label in choices}
    if raw not in allowed:
        raise ValidationError({field: f"Choose one of: {', '.join(allowed)}."})
    return raw


def _status(value: object) -> str:
    return _choice(value, field="status", choices=SourcingEntityStatus.choices, default=SourcingEntityStatus.ACTIVE)


def _availability(value: object) -> str:
    return _choice(value, field="availability", choices=SourcingAvailability.choices, default=SourcingAvailability.UNKNOWN)


def _rate_basis(value: object) -> str:
    return _choice(value, field="rate_basis", choices=SourcingRateBasis.choices, default=SourcingRateBasis.MONTH)


def _validation_message(exc: Exception) -> str:
    if isinstance(exc, ValidationError):
        if hasattr(exc, "message_dict"):
            parts = []
            for field, messages in exc.message_dict.items():
                values = messages if isinstance(messages, (list, tuple)) else [messages]
                parts.extend(f"{field}: {message}" for message in values)
            return "; ".join(parts)
        return "; ".join(exc.messages)
    return str(exc)


def read_import_rows(uploaded_file) -> tuple[list[str], list[dict[str, object]]]:
    name = Path(getattr(uploaded_file, "name", "upload")).name
    suffix = Path(name).suffix.casefold()
    size = getattr(uploaded_file, "size", None)
    if size is not None and size > MAX_IMPORT_BYTES:
        raise SourcingImportError(f"Import file exceeds {MAX_IMPORT_BYTES // (1024 * 1024)} MB.")
    raw = uploaded_file.read(MAX_IMPORT_BYTES + 1)
    if len(raw) > MAX_IMPORT_BYTES:
        raise SourcingImportError(f"Import file exceeds {MAX_IMPORT_BYTES // (1024 * 1024)} MB.")
    rows: list[list[object]] = []
    if suffix == ".csv":
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise SourcingImportError("CSV must be UTF-8 encoded.") from exc
        rows = [list(row) for row in csv.reader(io.StringIO(text))]
    elif suffix == ".xlsx":
        try:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            sheet = workbook.active
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            workbook.close()
        except Exception as exc:
            raise SourcingImportError("The XLSX workbook could not be read safely.") from exc
    else:
        raise SourcingImportError("Upload a .csv or .xlsx file.")
    if not rows:
        raise SourcingImportError("Import file is empty.")
    headers = [_header(value) for value in rows[0]]
    if not any(headers):
        raise SourcingImportError("Import header row is empty.")
    if len(headers) != len(set(headers)):
        raise SourcingImportError("Import contains duplicate column headers.")
    records = []
    for raw_row in rows[1:]:
        values = list(raw_row) + [None] * max(0, len(headers) - len(raw_row))
        record = {headers[index]: values[index] for index in range(len(headers)) if headers[index]}
        if any(_text(value) for value in record.values()):
            records.append(record)
            if len(records) > MAX_IMPORT_ROWS:
                raise SourcingImportError(f"Import is limited to {MAX_IMPORT_ROWS:,} data rows.")
    if not records:
        raise SourcingImportError("Import contains no data rows.")
    return headers, records


def _identity(dataset: str, row: dict[str, object]) -> str:
    if dataset in {"vendors", "materials", "manpower_suppliers", "trades", "subcontractors", "services"}:
        return _text(row.get("code")).upper() or "(missing code)"
    if dataset == "vendor_catalog":
        return f"{_text(row.get('vendor_code')).upper()} / {_text(row.get('material_code')).upper()}"
    if dataset == "subcontractor_service_catalog":
        return " / ".join((
            _text(row.get("subcontractor_code")).upper(),
            _text(row.get("service_code")).upper(),
            _text(row.get("scope_description")).casefold(),
            _text(row.get("service_location")).casefold(),
        ))
    return f"{_text(row.get('supplier_code')).upper()} / {_text(row.get('trade_code')).upper()}"


def _require_headers(dataset: str, headers: Iterable[str]) -> None:
    missing = REQUIRED_HEADERS[dataset] - set(headers)
    if missing:
        raise SourcingImportError(f"Missing required column(s): {', '.join(sorted(missing))}.")


def _upsert_vendor(*, actor_membership, row, request):
    company = actor_membership.company
    code = _text(row.get("code")).upper()
    required_values = {
        "name": "Vendor name",
        "display_name": "Display name",
        "mobile": "Primary contact mobile number",
        "email": "Primary contact email address",
        "address": "Address",
        "cr_number": "CR number",
        "vat_number": "VAT number",
    }
    missing_values = [label for field, label in required_values.items() if not _text(row.get(field))]
    if missing_values:
        raise ValidationError({"vendor": f"Required value(s) missing: {', '.join(missing_values)}."})
    existing = SourcingVendor.objects.for_company(company).filter(code=code).first()
    data = {
        "code": code,
        "name": _text(row.get("name")),
        "display_name": _text(row.get("display_name")),
        "cr_number": _text(row.get("cr_number")),
        "vat_number": _text(row.get("vat_number")),
        "company_phone": _text(row.get("company_phone") or row.get("phone")),
        "company_email": _text(row.get("company_email")),
        "website": _text(row.get("website")),
        "primary_contact_name": _text(row.get("primary_contact_name")),
        "mobile": _text(row.get("mobile")),
        "email": _text(row.get("email")),
        "address": _text(row.get("address")),
        "street_number": _text(row.get("street_number")),
        "district": _text(row.get("district")),
        "city": _text(row.get("city")),
        "region": _text(row.get("region")),
        "postal_code": _text(row.get("postal_code")),
        "notes": str(row.get("notes") or "").strip(),
    }
    obj = update_vendor(actor_membership=actor_membership, vendor_id=existing.pk, cleaned_data=data, request=request) if existing else create_vendor(actor_membership=actor_membership, cleaned_data=data, request=request)
    status = _status(row.get("status"))
    if obj.status != status:
        obj = set_vendor_status(actor_membership=actor_membership, vendor_id=obj.pk, status=status, request=request)
    return "update" if existing else "create", obj


def _upsert_material(*, actor_membership, row, request):
    company = actor_membership.company
    code = _text(row.get("code")).upper()
    existing = SourcingMaterial.objects.for_company(company).filter(code=code).first()
    data = {
        "code": code,
        "name": _text(row.get("name")),
        "category": _text(row.get("category")),
        "default_unit": _text(row.get("default_unit")),
        "aliases": _aliases(row.get("aliases")),
        "notes": str(row.get("notes") or "").strip(),
    }
    obj = update_material(actor_membership=actor_membership, material_id=existing.pk, cleaned_data=data, request=request) if existing else create_material(actor_membership=actor_membership, cleaned_data=data, request=request)
    active = _bool(row.get("is_active"), default=True)
    if obj.is_active != active:
        obj = set_material_active(actor_membership=actor_membership, material_id=obj.pk, is_active=active, request=request)
    return "update" if existing else "create", obj


def _upsert_vendor_catalog(*, actor_membership, row, request):
    company = actor_membership.company
    vendor_code = _text(row.get("vendor_code")).upper()
    material_code = _text(row.get("material_code")).upper()
    vendor = SourcingVendor.objects.for_company(company).filter(code=vendor_code, deleted_at__isnull=True).first()
    material = SourcingMaterial.objects.for_company(company).filter(code=material_code).first()
    if vendor is None:
        raise ValidationError({"vendor_code": f"Unknown Sourcing Vendor {vendor_code!r}."})
    if material is None:
        raise ValidationError({"material_code": f"Unknown Sourcing Material {material_code!r}."})
    existing = SourcingVendorOffer.objects.for_company(company).filter(vendor=vendor, material=material).first()
    data = {
        "material": material,
        "specification": _text(row.get("specification")),
        "brand": _text(row.get("brand")),
        "model": _text(row.get("model")),
        "available_quantity": _decimal(row.get("available_quantity"), field="available_quantity"),
        "unit": _text(row.get("unit")) or material.default_unit,
        "minimum_quantity": _decimal(row.get("minimum_quantity"), field="minimum_quantity"),
        "availability": _availability(row.get("availability")),
        "rate": _decimal(row.get("rate"), field="rate"),
        "currency": _text(row.get("currency")).upper() or "SAR",
        "rate_valid_until": _date(row.get("rate_valid_until"), field="rate_valid_until"),
        "lead_time": _text(row.get("lead_time")),
        "verification_note": _text(row.get("verification_note"))[:500],
        "notes": str(row.get("notes") or "").strip(),
    }
    verified_now = _bool(row.get("verified_now"), default=False)
    contact_name = _text(row.get("contact_name"))[:160]
    obj = (
        update_vendor_offer(actor_membership=actor_membership, vendor_id=vendor.pk, offer_id=existing.pk, cleaned_data=data, verified_now=verified_now, contact_name=contact_name, request=request)
        if existing else
        create_vendor_offer(actor_membership=actor_membership, vendor_id=vendor.pk, cleaned_data=data, verified_now=verified_now, contact_name=contact_name, request=request)
    )
    active = _bool(row.get("is_active"), default=True)
    if obj.is_active != active:
        obj = set_vendor_offer_active(actor_membership=actor_membership, vendor_id=vendor.pk, offer_id=obj.pk, is_active=active, request=request)
    return "update" if existing else "create", obj


def _upsert_manpower_supplier(*, actor_membership, row, request):
    company = actor_membership.company
    code = _text(row.get("code")).upper()
    existing = SourcingManpowerSupplier.objects.for_company(company).filter(code=code).first()
    data = {
        "code": code,
        "name": _text(row.get("name")),
        "primary_contact_name": _text(row.get("primary_contact_name")),
        "phone": _text(row.get("phone")),
        "mobile": _text(row.get("mobile")),
        "email": _text(row.get("email")),
        "city": _text(row.get("city")),
        "region": _text(row.get("region")),
        "cr_number": _text(row.get("cr_number")),
        "vat_number": _text(row.get("vat_number")),
        "notes": str(row.get("notes") or "").strip(),
    }
    obj = update_manpower_supplier(actor_membership=actor_membership, supplier_id=existing.pk, cleaned_data=data, request=request) if existing else create_manpower_supplier(actor_membership=actor_membership, cleaned_data=data, request=request)
    status = _status(row.get("status"))
    if obj.status != status:
        obj = set_manpower_supplier_status(actor_membership=actor_membership, supplier_id=obj.pk, status=status, request=request)
    return "update" if existing else "create", obj


def _upsert_trade(*, actor_membership, row, request):
    company = actor_membership.company
    code = _text(row.get("code")).upper()
    existing = SourcingTrade.objects.for_company(company).filter(code=code).first()
    data = {
        "code": code,
        "name": _text(row.get("name")),
        "category": _text(row.get("category")),
        "aliases": _aliases(row.get("aliases")),
        "notes": str(row.get("notes") or "").strip(),
    }
    obj = update_trade(actor_membership=actor_membership, trade_id=existing.pk, cleaned_data=data, request=request) if existing else create_trade(actor_membership=actor_membership, cleaned_data=data, request=request)
    active = _bool(row.get("is_active"), default=True)
    if obj.is_active != active:
        obj = set_trade_active(actor_membership=actor_membership, trade_id=obj.pk, is_active=active, request=request)
    return "update" if existing else "create", obj


def _upsert_workforce_catalog(*, actor_membership, row, request):
    company = actor_membership.company
    supplier_code = _text(row.get("supplier_code")).upper()
    trade_code = _text(row.get("trade_code")).upper()
    supplier = SourcingManpowerSupplier.objects.for_company(company).filter(code=supplier_code, deleted_at__isnull=True).first()
    trade = SourcingTrade.objects.for_company(company).filter(code=trade_code).first()
    if supplier is None:
        raise ValidationError({"supplier_code": f"Unknown Sourcing Manpower Supplier {supplier_code!r}."})
    if trade is None:
        raise ValidationError({"trade_code": f"Unknown Sourcing Trade {trade_code!r}."})
    existing = SourcingWorkforceOffer.objects.for_company(company).filter(supplier=supplier, trade=trade).first()
    data = {
        "trade": trade,
        "availability": _availability(row.get("availability")),
        "available_quantity": _integer(row.get("available_quantity"), field="available_quantity"),
        "rate": _decimal(row.get("rate"), field="rate"),
        "currency": _text(row.get("currency")).upper() or "SAR",
        "rate_basis": _rate_basis(row.get("rate_basis")),
        "overtime_rate": _decimal(row.get("overtime_rate"), field="overtime_rate"),
        "rate_valid_until": _date(row.get("rate_valid_until"), field="rate_valid_until"),
        "mobilization_lead_time": _text(row.get("mobilization_lead_time")),
        "work_location": _text(row.get("work_location")),
        "verification_note": _text(row.get("verification_note"))[:500],
        "notes": str(row.get("notes") or "").strip(),
    }
    verified_now = _bool(row.get("verified_now"), default=False)
    contact_name = _text(row.get("contact_name"))[:160]
    obj = (
        update_workforce_offer(actor_membership=actor_membership, supplier_id=supplier.pk, offer_id=existing.pk, cleaned_data=data, verified_now=verified_now, contact_name=contact_name, request=request)
        if existing else
        create_workforce_offer(actor_membership=actor_membership, supplier_id=supplier.pk, cleaned_data=data, verified_now=verified_now, contact_name=contact_name, request=request)
    )
    active = _bool(row.get("is_active"), default=True)
    if obj.is_active != active:
        obj = set_workforce_offer_active(actor_membership=actor_membership, supplier_id=supplier.pk, offer_id=obj.pk, is_active=active, request=request)
    return "update" if existing else "create", obj



def _upsert_subcontractor(*, actor_membership, row, request):
    company = actor_membership.company
    code = _text(row.get("code")).upper()
    obj = SourcingSubcontractor.objects.for_company(company).filter(code=code).first()
    if obj is not None and (obj.deleted_at or obj.archived_at):
        raise ValidationError("Restore this archived/trashed subcontractor before importing updates.")
    fields = (
        "name", "display_name", "primary_contact_name", "company_phone", "company_email",
        "mobile", "email", "website", "address", "street_number", "district",
        "city", "region", "postal_code", "service_coverage", "cr_number", "vat_number",
    )
    data = {key: _text(row.get(key)) for key in fields}
    data.update(code=code, notes=str(row.get("notes") or "").strip())
    subcontractor = (
        update_subcontractor(actor_membership=actor_membership, subcontractor_id=obj.pk, cleaned_data=data, request=request)
        if obj else create_subcontractor(actor_membership=actor_membership, cleaned_data=data, request=request)
    )
    status = _status(row.get("status"))
    if subcontractor.status != status:
        set_subcontractor_status(actor_membership=actor_membership, subcontractor_id=subcontractor.pk,
                                 status=status, request=request)
    # last_verified_at from exported files is intentionally ignored. Company
    # verification must be explicitly confirmed with independent evidence.
    return ("update" if obj else "create"), subcontractor


def _upsert_service(*, actor_membership, row, request):
    company = actor_membership.company
    code = _text(row.get("code")).upper()
    existing = SourcingService.objects.for_company(company).filter(code=code).first()
    basis = _text(row.get("default_rate_basis")).casefold()
    if basis and basis not in SourcingServiceRateBasis.values:
        raise ValidationError({"default_rate_basis": "Choose a valid service rate basis."})
    data = {
        "code": code,
        "name": _text(row.get("name")),
        "category": _text(row.get("category")),
        "default_rate_basis": basis,
        "aliases": _aliases(row.get("aliases")),
        "description": str(row.get("description") or "").strip(),
        "notes": str(row.get("notes") or "").strip(),
    }
    service = (
        update_service(actor_membership=actor_membership, service_id=existing.pk, cleaned_data=data, request=request)
        if existing else create_service(actor_membership=actor_membership, cleaned_data=data, request=request)
    )
    active = _bool(row.get("is_active"), default=True)
    if service.is_active != active:
        set_service_active(actor_membership=actor_membership, service_id=service.pk, is_active=active, request=request)
    return ("update" if existing else "create"), service


def _upsert_subcontractor_service_catalog(*, actor_membership, row, request):
    company = actor_membership.company
    subcontractor_code = _text(row.get("subcontractor_code")).upper()
    service_code = _text(row.get("service_code")).upper()
    subcontractor = SourcingSubcontractor.objects.for_company(company).filter(
        code=subcontractor_code, deleted_at__isnull=True, archived_at__isnull=True,
    ).first()
    service = SourcingService.objects.for_company(company).filter(code=service_code).first()
    if subcontractor is None:
        raise ValidationError({"subcontractor_code": "Unknown or archived/trashed Subcontractor in this company."})
    if service is None:
        raise ValidationError({"service_code": "Unknown Service in this company."})
    scope = _text(row.get("scope_description"))
    location = _text(row.get("service_location"))
    existing = SourcingSubcontractorService.objects.for_company(company).filter(
        subcontractor=subcontractor, service=service,
        scope_description__iexact=scope, service_location__iexact=location,
    ).first()
    rate_basis = _choice(row.get("rate_basis"), field="rate_basis",
                         choices=SourcingServiceRateBasis.choices,
                         default=SourcingServiceRateBasis.QUOTATION)
    data = {
        "service": service,
        "scope_description": scope,
        "service_location": location,
        "availability": _availability(row.get("availability")),
        "rate": _service_rate(row.get("rate")),
        "currency": _text(row.get("currency")).upper() or "SAR",
        "rate_basis": rate_basis,
        "rate_valid_until": _date(row.get("rate_valid_until"), field="rate_valid_until"),
        "lead_time": _text(row.get("lead_time")),
        "verification_note": _text(row.get("verification_note"))[:500],
        "notes": str(row.get("notes") or "").strip(),
    }
    verified_now = _bool(row.get("verified_now"), default=False)
    contact_name = _text(row.get("contact_name"))[:160]
    offer = (
        update_service_offer(actor_membership=actor_membership, subcontractor_id=subcontractor.pk,
                             offer_id=existing.pk, cleaned_data=data, verified_now=verified_now,
                             contact_name=contact_name, request=request)
        if existing else
        create_service_offer(actor_membership=actor_membership, subcontractor_id=subcontractor.pk,
                             cleaned_data=data, verified_now=verified_now,
                             contact_name=contact_name, request=request)
    )
    active = _bool(row.get("is_active"), default=True)
    if offer.is_active != active:
        set_service_offer_active(actor_membership=actor_membership, subcontractor_id=subcontractor.pk,
                                 offer_id=offer.pk, is_active=active, request=request)
    # Exported last_verified_at is read-only; only verified_now creates evidence.
    return ("update" if existing else "create"), offer

_UPSERT = {
    "vendors": _upsert_vendor,
    "materials": _upsert_material,
    "vendor_catalog": _upsert_vendor_catalog,
    "manpower_suppliers": _upsert_manpower_supplier,
    "trades": _upsert_trade,
    "workforce_catalog": _upsert_workforce_catalog,
    "subcontractors": _upsert_subcontractor,
    "services": _upsert_service,
    "subcontractor_service_catalog": _upsert_subcontractor_service_catalog,
}


def import_sourcing_rows(*, actor_membership, dataset: str, headers: list[str], records: list[dict[str, object]], dry_run: bool, request=None) -> ImportResult:
    if dataset not in _UPSERT:
        raise SourcingImportError("Unsupported Sourcing import dataset.")
    _require_headers(dataset, headers)
    results: list[ImportRowResult] = []
    seen = set()
    with transaction.atomic():
        for offset, row in enumerate(records, start=2):
            identity = _identity(dataset, row)
            key = identity.casefold()
            if key in seen:
                results.append(ImportRowResult(offset, "error", identity, "Duplicate identity appears more than once in this import."))
                continue
            seen.add(key)
            try:
                with transaction.atomic():
                    outcome, _obj = _UPSERT[dataset](actor_membership=actor_membership, row=row, request=request)
                results.append(ImportRowResult(offset, outcome, identity))
            except (ValidationError, IntegrityError) as exc:
                results.append(ImportRowResult(offset, "error", identity, _validation_message(exc)))
        has_errors = any(row.outcome == "error" for row in results)
        if dry_run or has_errors:
            transaction.set_rollback(True)
    result = ImportResult(dataset=dataset, dry_run=bool(dry_run), rows=tuple(results))
    if not dry_run and not result.errors:
        record_audit_event(
            company=actor_membership.company,
            area=AuditArea.SOURCING,
            action="sourcing.data_import.completed",
            object_type="sourcing.SourcingDataExchange",
            object_id=actor_membership.company_id,
            object_label=f"Sourcing import · {dataset}",
            actor_membership=actor_membership,
            metadata={"dataset": dataset, "rows": result.total, "created": result.creates, "updated": result.updates, "allOrNothing": True},
            request=request,
        )
    return result


def _safe_export_cell(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        if isinstance(value, datetime) and timezone.is_aware(value):
            value = timezone.localtime(value)
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, bool):
        return "yes" if value else "no"
    value = str(value)
    if value.startswith(_FORMULA_PREFIXES) or value.lstrip().startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def _contains(value: str, query: str) -> bool:
    return query.casefold() in value.casefold()


def _query_decimal(value: object) -> Decimal | None:
    try:
        raw = _text(value)
        if not raw:
            return None
        parsed = Decimal(raw)
        return (parsed if parsed.is_finite() and parsed >= 0
                and parsed < Decimal("100000000000000")
                and parsed.as_tuple().exponent >= -4 else None)
    except (InvalidOperation, ValueError):
        return None


def _apply_freshness_filter(qs, *, company, value: str):
    value = _text(value).casefold()
    if value not in {FRESHNESS_FRESH, FRESHNESS_NEEDS_VERIFICATION, FRESHNESS_STALE, FRESHNESS_NEVER}:
        return qs
    fresh_cutoff, stale_cutoff = sourcing_freshness_policy(company=company).cutoffs(now=timezone.now())
    if value == FRESHNESS_FRESH:
        return qs.filter(last_verified_at__gte=fresh_cutoff)
    if value == FRESHNESS_NEEDS_VERIFICATION:
        return qs.filter(last_verified_at__lt=fresh_cutoff, last_verified_at__gte=stale_cutoff)
    if value == FRESHNESS_STALE:
        return qs.filter(last_verified_at__lt=stale_cutoff)
    return qs.filter(last_verified_at__isnull=True)


def _directory_export_search(qs, *, company, dataset: str, query: str):
    """Match capability and contact searches on company exports just like the directory."""
    if not query:
        return qs
    normalized = normalize_text(query)
    if dataset == "vendors":
        contacts = SourcingVendorContact.objects.for_company(company).filter(
            vendor_id=OuterRef("pk"), is_active=True,
        ).filter(Q(first_name__icontains=query) | Q(last_name__icontains=query)
                 | Q(email__icontains=query) | Q(mobile__icontains=query)
                 | Q(work_phone__icontains=query))
        offers = SourcingVendorOffer.objects.for_company(company).filter(
            vendor_id=OuterRef("pk"), is_active=True, material__is_active=True,
        ).filter(Q(material__code__icontains=query) | Q(material__name__icontains=query)
                 | Q(material__category__icontains=query)
                 | Q(material__normalized_aliases__icontains=normalized)
                 | Q(specification__icontains=query) | Q(brand__icontains=query)
                 | Q(model__icontains=query))
    elif dataset == "manpower_suppliers":
        contacts = SourcingManpowerContact.objects.for_company(company).filter(
            supplier_id=OuterRef("pk"), is_active=True,
        ).filter(Q(first_name__icontains=query) | Q(last_name__icontains=query)
                 | Q(email__icontains=query) | Q(mobile__icontains=query)
                 | Q(work_phone__icontains=query))
        offers = SourcingWorkforceOffer.objects.for_company(company).filter(
            supplier_id=OuterRef("pk"), is_active=True, trade__is_active=True,
        ).filter(Q(trade__code__icontains=query) | Q(trade__name__icontains=query)
                 | Q(trade__category__icontains=query)
                 | Q(trade__normalized_aliases__icontains=normalized)
                 | Q(work_location__icontains=query) | Q(mobilization_lead_time__icontains=query))
    else:
        contacts = SourcingSubcontractorContact.objects.for_company(company).filter(
            subcontractor_id=OuterRef("pk"),
        ).filter(Q(first_name__icontains=query) | Q(last_name__icontains=query)
                 | Q(email__icontains=query) | Q(mobile__icontains=query))
        offers = SourcingSubcontractorService.objects.for_company(company).filter(
            subcontractor_id=OuterRef("pk"), is_active=True, service__is_active=True,
        ).filter(Q(service__code__icontains=query) | Q(service__name__icontains=query)
                 | Q(service__category__icontains=query)
                 | Q(service__normalized_aliases__icontains=normalized)
                 | Q(scope_description__icontains=query) | Q(service_location__icontains=query))
    return qs.annotate(_export_contact_hit=Exists(contacts), _export_offer_hit=Exists(offers))


def export_rows(*, company, dataset: str, params) -> list[list[object]]:
    q = _text(params.get("q"))
    rows: list[list[object]] = []
    if dataset == "vendors":
        qs = SourcingVendor.objects.for_company(company)
        status = _text(params.get("status")).casefold()
        if status == "trash":
            qs = qs.filter(deleted_at__isnull=False, purge_after__gt=timezone.now())
        else:
            qs = qs.filter(deleted_at__isnull=True)
            if status in {SourcingEntityStatus.ACTIVE, SourcingEntityStatus.INACTIVE}:
                qs = qs.filter(archived_at__isnull=True, status=status)
            elif status == "archived":
                qs = qs.filter(archived_at__isnull=False)
        if q:
            qs = _directory_export_search(qs, company=company, dataset=dataset, query=q)
            qs = qs.filter(
                Q(code__icontains=q)
                | Q(name__icontains=q)
                | Q(display_name__icontains=q)
                | Q(company_phone__icontains=q)
                | Q(company_email__icontains=q)
                | Q(primary_contact_name__icontains=q)
                | Q(mobile__icontains=q)
                | Q(email__icontains=q)
                | Q(address__icontains=q)
                | Q(street_number__icontains=q)
                | Q(district__icontains=q)
                | Q(city__icontains=q)
                | Q(region__icontains=q)
                | Q(postal_code__icontains=q)
                | Q(cr_number__icontains=q)
                | Q(vat_number__icontains=q)
                | Q(_export_contact_hit=True) | Q(_export_offer_hit=True)
            )
        if params.get("unified") == "1":
            # Use EXACTLY the tenant-scoped, company-first search queryset.
            from .selectors.vendors import vendor_directory_page
            qs = vendor_directory_page(company=company, params=params).page_obj.paginator.object_list
        for obj in qs.order_by("name", "code").iterator(chunk_size=500):
            rows.append([
                obj.code, obj.name, obj.display_name, obj.cr_number, obj.vat_number,
                obj.company_phone, obj.company_email, obj.website, obj.primary_contact_name,
                obj.mobile, obj.email, obj.address, obj.street_number, obj.district, obj.city,
                obj.region, obj.postal_code, obj.status, obj.notes, obj.last_verified_at,
            ])
    elif dataset == "materials":
        qs = SourcingMaterial.objects.for_company(company)
        status = _text(params.get("status")).casefold()
        if status == "active":
            qs = qs.filter(is_active=True)
        elif status == "inactive":
            qs = qs.filter(is_active=False)
        category = _text(params.get("category"))
        if category:
            qs = qs.filter(category__iexact=category)
        if q:
            qs = qs.filter(Q(code__icontains=q) | Q(name__icontains=q) | Q(category__icontains=q) | Q(normalized_aliases__icontains=q))
        for obj in qs.order_by("category", "name", "code").iterator(chunk_size=500):
            rows.append([obj.code, obj.name, obj.category, obj.default_unit, " | ".join(obj.aliases or []), obj.is_active, obj.notes])
    elif dataset == "vendor_catalog":
        qs = SourcingVendorOffer.objects.for_company(company).select_related("vendor", "material", "verified_by").filter(vendor__deleted_at__isnull=True)
        finder_mode = any(key in params for key in ("freshness", "category", "vendor", "location", "rate_min", "rate_max", "page_size", "sort", "dir"))
        if finder_mode or params.get("view") == "offers":
            qs = qs.filter(is_active=True, material__is_active=True, vendor__archived_at__isnull=True, vendor__status=SourcingEntityStatus.ACTIVE)
        vendor_code = _text(params.get("vendor_code"))
        material_code = _text(params.get("material_code"))
        vendor_filter = _text(params.get("vendor"))
        category = _text(params.get("category"))
        location = _text(params.get("location"))
        availability = _text(params.get("availability"))
        rate_min = _query_decimal(params.get("rate_min"))
        rate_max = _query_decimal(params.get("rate_max"))
        if vendor_code:
            qs = qs.filter(vendor__code__iexact=vendor_code)
        if material_code:
            qs = qs.filter(material__code__iexact=material_code)
        if vendor_filter:
            qs = qs.filter(Q(vendor__code__icontains=vendor_filter) | Q(vendor__name__icontains=vendor_filter) | Q(vendor__display_name__icontains=vendor_filter))
        if category:
            qs = qs.filter(material__category__iexact=category)
        if location:
            qs = qs.filter(Q(vendor__address__icontains=location) | Q(vendor__district__icontains=location) | Q(vendor__city__icontains=location) | Q(vendor__region__icontains=location) | Q(vendor__postal_code__icontains=location))
        if availability in {value for value, _ in SourcingAvailability.choices}:
            qs = qs.filter(availability=availability)
        if rate_min is not None:
            qs = qs.filter(rate__gte=rate_min)
        if rate_max is not None:
            qs = qs.filter(rate__lte=rate_max)
        qs = _apply_freshness_filter(qs, company=company, value=params.get("freshness"))
        if q:
            qs = qs.filter(Q(vendor__name__icontains=q) | Q(vendor__code__icontains=q) | Q(material__name__icontains=q) | Q(material__code__icontains=q) | Q(material__category__icontains=q) | Q(material__normalized_aliases__icontains=normalize_text(q)) | Q(specification__icontains=q) | Q(brand__icontains=q) | Q(model__icontains=q))
        if params.get("unified") == "1":
            from .selectors.unified import capability_queryset, parse_filters
            from .selectors.vendors import vendor_directory_page
            matched_companies = vendor_directory_page(company=company, params=params).page_obj.paginator.object_list.order_by().values('pk')
            qs = capability_queryset(company=company, kind='vendor', filters=parse_filters(kind='vendor', params=params), query=q)
            qs = qs.filter(vendor_id__in=matched_companies).select_related('vendor', 'material')
        for obj in qs.order_by("material__name", "vendor__name").iterator(chunk_size=500):
            rows.append([obj.vendor.code, obj.material.code, obj.specification, obj.brand, obj.model, obj.availability, obj.available_quantity, obj.unit, obj.minimum_quantity, obj.rate, obj.currency, obj.rate_valid_until, obj.lead_time, "", "", obj.verification_note, obj.is_active, obj.notes, obj.last_verified_at])
    elif dataset == "manpower_suppliers":
        qs = SourcingManpowerSupplier.objects.for_company(company)
        status = _text(params.get("status")).casefold()
        if status == "trash":
            qs = qs.filter(deleted_at__isnull=False, purge_after__gt=timezone.now())
        else:
            qs = qs.filter(deleted_at__isnull=True)
            if status in {SourcingEntityStatus.ACTIVE, SourcingEntityStatus.INACTIVE}:
                qs = qs.filter(archived_at__isnull=True, status=status)
            elif status == "archived":
                qs = qs.filter(archived_at__isnull=False)
        if q:
            qs = _directory_export_search(qs, company=company, dataset=dataset, query=q)
            qs = qs.filter(Q(code__icontains=q) | Q(name__icontains=q) | Q(primary_contact_name__icontains=q) | Q(phone__icontains=q) | Q(email__icontains=q) | Q(city__icontains=q) | Q(region__icontains=q) | Q(cr_number__icontains=q) | Q(vat_number__icontains=q)
                           | Q(_export_contact_hit=True) | Q(_export_offer_hit=True))
        if params.get("unified") == "1":
            # Use EXACTLY the tenant-scoped, company-first search queryset.
            from .selectors.manpower import manpower_directory_page
            qs = manpower_directory_page(company=company, params=params).page_obj.paginator.object_list
        for obj in qs.order_by("name", "code").iterator(chunk_size=500):
            rows.append([obj.code, obj.name, obj.primary_contact_name, obj.phone, obj.mobile, obj.email, obj.city, obj.region, obj.cr_number, obj.vat_number, obj.status, obj.notes, obj.last_verified_at])
    elif dataset == "trades":
        qs = SourcingTrade.objects.for_company(company)
        status = _text(params.get("status")).casefold()
        if status == "active":
            qs = qs.filter(is_active=True)
        elif status == "inactive":
            qs = qs.filter(is_active=False)
        category = _text(params.get("category"))
        if category:
            qs = qs.filter(category__iexact=category)
        if q:
            qs = qs.filter(Q(code__icontains=q) | Q(name__icontains=q) | Q(category__icontains=q) | Q(normalized_aliases__icontains=q))
        for obj in qs.order_by("category", "name", "code").iterator(chunk_size=500):
            rows.append([obj.code, obj.name, obj.category, " | ".join(obj.aliases or []), obj.is_active, obj.notes])
    elif dataset == "workforce_catalog":
        qs = SourcingWorkforceOffer.objects.for_company(company).select_related("supplier", "trade", "verified_by").filter(supplier__deleted_at__isnull=True)
        finder_mode = any(key in params for key in ("freshness", "category", "supplier", "location", "rate_min", "rate_max", "page_size", "sort", "dir"))
        if finder_mode or params.get("view") == "offers":
            qs = qs.filter(is_active=True, trade__is_active=True, supplier__archived_at__isnull=True, supplier__status=SourcingEntityStatus.ACTIVE)
        supplier_code = _text(params.get("supplier_code"))
        trade_code = _text(params.get("trade_code"))
        supplier_filter = _text(params.get("supplier"))
        category = _text(params.get("category"))
        location = _text(params.get("location"))
        availability = _text(params.get("availability"))
        rate_basis = _text(params.get("rate_basis"))
        rate_min = _query_decimal(params.get("rate_min"))
        rate_max = _query_decimal(params.get("rate_max"))
        if supplier_code:
            qs = qs.filter(supplier__code__iexact=supplier_code)
        if trade_code:
            qs = qs.filter(trade__code__iexact=trade_code)
        if supplier_filter:
            qs = qs.filter(Q(supplier__code__icontains=supplier_filter) | Q(supplier__name__icontains=supplier_filter))
        if category:
            qs = qs.filter(trade__category__iexact=category)
        if location:
            qs = qs.filter(Q(work_location__icontains=location) | Q(supplier__city__icontains=location) | Q(supplier__region__icontains=location))
        if availability in {value for value, _ in SourcingAvailability.choices}:
            qs = qs.filter(availability=availability)
        if rate_basis in {value for value, _ in SourcingRateBasis.choices}:
            qs = qs.filter(rate_basis=rate_basis)
        if rate_min is not None:
            qs = qs.filter(rate__gte=rate_min)
        if rate_max is not None:
            qs = qs.filter(rate__lte=rate_max)
        qs = _apply_freshness_filter(qs, company=company, value=params.get("freshness"))
        if q:
            qs = qs.filter(Q(supplier__name__icontains=q) | Q(supplier__code__icontains=q) | Q(trade__name__icontains=q) | Q(trade__code__icontains=q) | Q(trade__normalized_aliases__icontains=normalize_text(q)) | Q(trade__category__icontains=q) | Q(work_location__icontains=q) | Q(mobilization_lead_time__icontains=q))
        if params.get("unified") == "1":
            from .selectors.unified import capability_queryset, parse_filters
            from .selectors.manpower import manpower_directory_page
            matched_companies = manpower_directory_page(company=company, params=params).page_obj.paginator.object_list.order_by().values('pk')
            qs = capability_queryset(company=company, kind='manpower', filters=parse_filters(kind='manpower', params=params), query=q)
            qs = qs.filter(supplier_id__in=matched_companies).select_related('supplier', 'trade')
        for obj in qs.order_by("trade__name", "supplier__name").iterator(chunk_size=500):
            rows.append([obj.supplier.code, obj.trade.code, obj.availability, obj.available_quantity, obj.rate, obj.currency, obj.rate_basis, obj.overtime_rate, obj.rate_valid_until, obj.mobilization_lead_time, obj.work_location, "", "", obj.verification_note, obj.is_active, obj.notes, obj.last_verified_at])
    elif dataset == "subcontractors":
        qs = SourcingSubcontractor.objects.for_company(company)
        status = _text(params.get("status")).casefold()
        if status == "trash":
            qs = qs.filter(deleted_at__isnull=False, purge_after__gt=timezone.now())
        else:
            qs = qs.filter(deleted_at__isnull=True)
            if status in {SourcingEntityStatus.ACTIVE, SourcingEntityStatus.INACTIVE}:
                qs = qs.filter(archived_at__isnull=True, status=status)
            elif status == "archived":
                qs = qs.filter(archived_at__isnull=False)
        if q:
            qs = _directory_export_search(qs, company=company, dataset=dataset, query=q)
            qs = qs.filter(
                Q(code__icontains=q) | Q(name__icontains=q) | Q(display_name__icontains=q)
                | Q(primary_contact_name__icontains=q) | Q(company_email__icontains=q)
                | Q(email__icontains=q) | Q(city__icontains=q) | Q(region__icontains=q)
                | Q(cr_number__icontains=q) | Q(vat_number__icontains=q)
                | Q(service_coverage__icontains=q)
                | Q(_export_contact_hit=True) | Q(_export_offer_hit=True)
            )
        if params.get("unified") == "1":
            # Use EXACTLY the tenant-scoped, company-first search queryset.
            from .selectors.subcontractors import subcontractor_directory_page
            qs = subcontractor_directory_page(company=company, params=params).page_obj.paginator.object_list
        for obj in qs.order_by("name", "code").iterator(chunk_size=500):
            rows.append([
                obj.code, obj.name, obj.display_name, obj.primary_contact_name,
                obj.company_phone, obj.company_email, obj.mobile, obj.email,
                obj.website, obj.address, obj.street_number, obj.district,
                obj.city, obj.region, obj.postal_code, obj.service_coverage,
                obj.cr_number, obj.vat_number, obj.status, obj.notes, obj.last_verified_at,
            ])
    elif dataset == "services":
        qs = SourcingService.objects.for_company(company)
        status = _text(params.get("status")).casefold()
        if status == "active":
            qs = qs.filter(is_active=True)
        elif status == "inactive":
            qs = qs.filter(is_active=False)
        category = _text(params.get("category"))
        if category:
            qs = qs.filter(category__iexact=category)
        if q:
            qs = qs.filter(Q(code__icontains=q) | Q(name__icontains=q)
                           | Q(category__icontains=q) | Q(normalized_aliases__icontains=q))
        for obj in qs.order_by("category", "name", "code").iterator(chunk_size=500):
            rows.append([obj.code, obj.name, obj.category, obj.default_rate_basis,
                         " | ".join(obj.aliases or []), obj.description,
                         obj.is_active, obj.notes])
    elif dataset == "subcontractor_service_catalog":
        qs = SourcingSubcontractorService.objects.for_company(company).select_related(
            "subcontractor", "service"
        ).filter(subcontractor__deleted_at__isnull=True)
        finder_mode = any(key in params for key in (
            "freshness", "category", "subcontractor", "location", "currency",
            "rate_min", "rate_max", "page_size", "sort", "dir",
        ))
        if finder_mode or params.get("view") == "offers":
            qs = qs.filter(is_active=True, service__is_active=True,
                           subcontractor__archived_at__isnull=True,
                           subcontractor__status=SourcingEntityStatus.ACTIVE)
        subcontractor_code = _text(params.get("subcontractor_code"))
        service_code = _text(params.get("service_code"))
        subcontractor_filter = _text(params.get("subcontractor"))
        category = _text(params.get("category"))
        location = _text(params.get("location"))
        availability = _text(params.get("availability"))
        rate_basis = _text(params.get("rate_basis"))
        currency = _text(params.get("currency")).upper()
        rate_min = _query_decimal(params.get("rate_min"))
        rate_max = _query_decimal(params.get("rate_max"))
        if subcontractor_code:
            qs = qs.filter(subcontractor__code__iexact=subcontractor_code)
        if service_code:
            qs = qs.filter(service__code__iexact=service_code)
        if subcontractor_filter:
            qs = qs.filter(Q(subcontractor__name__icontains=subcontractor_filter)
                           | Q(subcontractor__code__icontains=subcontractor_filter)
                           | Q(subcontractor__display_name__icontains=subcontractor_filter))
        if category:
            qs = qs.filter(service__category__iexact=category)
        if location:
            qs = qs.filter(Q(service_location__icontains=location)
                           | Q(subcontractor__service_coverage__icontains=location)
                           | Q(subcontractor__city__icontains=location)
                           | Q(subcontractor__region__icontains=location))
        if availability in SourcingAvailability.values:
            qs = qs.filter(availability=availability)
        if rate_basis in SourcingServiceRateBasis.values:
            qs = qs.filter(rate_basis=rate_basis)
        if currency:
            qs = qs.filter(currency=currency)
        if rate_min is not None:
            qs = qs.filter(rate__gte=rate_min)
        if rate_max is not None:
            qs = qs.filter(rate__lte=rate_max)
        qs = _apply_freshness_filter(qs, company=company, value=params.get("freshness"))
        if q:
            qs = qs.filter(Q(service__name__icontains=q) | Q(service__code__icontains=q)
                           | Q(service__normalized_aliases__icontains=normalize_text(q))
                           | Q(service__category__icontains=q)
                           | Q(subcontractor__name__icontains=q) | Q(subcontractor__code__icontains=q)
                           | Q(scope_description__icontains=q) | Q(service_location__icontains=q))
        if params.get("unified") == "1":
            from .selectors.unified import capability_queryset, parse_filters
            from .selectors.subcontractors import subcontractor_directory_page
            matched_companies = subcontractor_directory_page(company=company, params=params).page_obj.paginator.object_list.order_by().values('pk')
            qs = capability_queryset(company=company, kind='subcontractor', filters=parse_filters(kind='subcontractor', params=params), query=q)
            qs = qs.filter(subcontractor_id__in=matched_companies).select_related('subcontractor', 'service')
        for obj in qs.order_by("service__name", "subcontractor__name", "scope_description", "service_location").iterator(chunk_size=500):
            rows.append([
                obj.subcontractor.code, obj.service.code, obj.scope_description,
                obj.service_location, obj.availability, obj.rate, obj.currency,
                obj.rate_basis, obj.rate_valid_until, obj.lead_time, "", "",
                obj.verification_note, obj.is_active, obj.notes, obj.last_verified_at,
            ])
    else:
        raise ValidationError("Unsupported Sourcing export dataset.")
    return rows


def build_export(*, actor_membership, dataset: str, fmt: str, params, request=None) -> tuple[bytes, str, str, int]:
    if dataset not in HEADERS:
        raise ValidationError("Unsupported Sourcing export dataset.")
    fmt = _text(fmt).casefold() or "xlsx"
    if fmt not in {"csv", "xlsx"}:
        raise ValidationError("Export format must be csv or xlsx.")
    rows = export_rows(company=actor_membership.company, dataset=dataset, params=params)
    headers = HEADERS[dataset]
    stamp = timezone.localdate().isoformat()
    filename = f"sescco-sourcing-{dataset}-{stamp}.{fmt}"
    if fmt == "csv":
        stream = io.StringIO(newline="")
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(headers)
        for row in rows:
            writer.writerow([_safe_export_cell(value) for value in row])
        payload = ("\ufeff" + stream.getvalue()).encode("utf-8")
        content_type = "text/csv; charset=utf-8"
    else:
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet(title="Sourcing")
        sheet.append(headers)
        for row in rows:
            sheet.append([_safe_export_cell(value) for value in row])
        stream = io.BytesIO()
        workbook.save(stream)
        payload = stream.getvalue()
        content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    record_audit_event(
        company=actor_membership.company,
        area=AuditArea.SOURCING,
        action="sourcing.data_exported",
        object_type="sourcing.SourcingDataExchange",
        object_id=actor_membership.company_id,
        object_label=f"Sourcing export · {dataset}",
        actor_membership=actor_membership,
        metadata={"dataset": dataset, "format": fmt, "rows": len(rows), "filtered": bool(params)},
        request=request,
    )
    return payload, filename, content_type, len(rows)


def build_import_template(*, dataset: str, fmt: str) -> tuple[bytes, str, str]:
    """Generate a header-only template; no tenant or supplier information leaks."""
    if dataset not in HEADERS:
        raise ValidationError("Unsupported Sourcing import dataset.")
    fmt = _text(fmt).casefold() or "xlsx"
    if fmt not in {"csv", "xlsx"}:
        raise ValidationError("Template format must be csv or xlsx.")
    headers = HEADERS[dataset]
    filename = f"sescco-sourcing-{dataset}-import-template.{fmt}"
    if fmt == "csv":
        stream = io.StringIO(newline="")
        csv.writer(stream, lineterminator="\n").writerow(headers)
        return ("\ufeff" + stream.getvalue()).encode("utf-8"), filename, "text/csv; charset=utf-8"
    workbook = Workbook(write_only=True)
    workbook.create_sheet(title="Sourcing Import").append(headers)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue(), filename, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
